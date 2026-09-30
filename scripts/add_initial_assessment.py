"""Add the Initial Assessment slide as Slide 1 to GCSE_Maths_Revision_Starters_prev4.pptx.
Preserves all existing 21 slides without modifying any of their content.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PPTX_PATH = ROOT / 'GCSE_Maths_Revision_Starters_prev4.pptx'
ANSWERS_PATH = ROOT / 'GCSE_Maths_Revision_Starters_answers_prev4.md'

BG, INK, MUTED, LINE = 'F7F8FA', '182B3A', '526472', 'D8E0E6'
COLORS = ['176B73', '3559A2', '754B87']
TINTS = ['EDF6F5', 'EFF3FA', 'F5F0F7']
MARGIN, FULL_W = 0.42, 15.16
COL_W = FULL_W / 3  # ~5.0533


def box(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    effect = shape._element.find('.//' + qn('a:effectRef'))
    if effect is not None:
        effect.set('idx', '0')
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(0.6)
    else:
        shape.line.fill.background()
    return shape


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False, valign=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.10)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = valign
    for i, line in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size * 1.08)
    return shape


def build_assessment_slide(prs):
    # Add a blank slide using blank layout (layout 6)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)

    # 3-segment top color bar reflecting the 3 modules
    box(s, 0, 0, 16 / 3, 0.10, COLORS[0])
    box(s, 16 / 3, 0, 16 / 3, 0.10, COLORS[1])
    box(s, 32 / 3, 0, 16 / 3, 0.10, COLORS[2])

    # Title & Subtitle
    text(s, MARGIN, 0.24, 13.6, 0.49, 'Initial Assessment', 29, bold=True)
    text(s, MARGIN, 0.78, 14.8, 0.30, 'Work through each question showing your full written method. No calculators.', 14, MUTED)

    # Footer
    text(s, MARGIN, 8.62, 11, 0.23, 'GCSE Maths revision  •  Record your working in your booklet or webapp.', 11, MUTED)
    text(s, 12.5, 8.60, 3.05, 0.25, 'Initial Assessment  •  12 Questions', 11, MUTED)

    # Column headers
    headers = ['Multiplication', 'Division', 'Application in Exam Questions']
    for c in range(3):
        x = MARGIN + c * COL_W
        box(s, x, 1.23, COL_W, 0.48, COLORS[c])
        text(s, x + 0.06, 1.26, COL_W - 0.12, 0.37, headers[c], 20, 'FFFFFF', True,
             valign=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

    # Questions definitions
    mult_q = ['43 × 6', '68 × 4', '476 × 82', '6,038 × 591']
    div_q = ['864 ÷ 4', '132 ÷ 8', '165 ÷ 6', '615 ÷ 12']
    app_q = [
        'A ribbon is 144 cm long and cut into 6 equal pieces. Find the length of each piece.',
        'A hall has 14 rows with 25 seats in each row. How many seats are there altogether?',
        '145 eggs are packed into boxes of 6. How many full boxes can be made?',
        'A coach hire company charges £75 per coach. 130 passengers need coaches with 24 seats each. Find the total hire cost.'
    ]
    all_cols = [mult_q, div_q, app_q]

    y_start = 1.85
    card_h = 1.48
    gap = 0.14

    for c in range(3):
        x = MARGIN + c * COL_W
        col_questions = all_cols[c]
        for r, q in enumerate(col_questions):
            y = y_start + r * (card_h + gap)
            card_fill = 'FFFFFF' if r % 2 == 0 else TINTS[c]
            box(s, x, y, COL_W, card_h, card_fill, LINE)

            # Number badge
            text(s, x + 0.12, y + 0.14, 0.45, 0.38, f'{r+1}.', 19, COLORS[c], True)

            # Question content
            if c < 2:
                text(s, x + 0.65, y + 0.38, COL_W - 0.85, 0.70, q, 30, bold=True,
                     valign=MSO_ANCHOR.MIDDLE)
            else:
                text(s, x + 0.58, y + 0.10, COL_W - 0.72, card_h - 0.20, q, 16, bold=False,
                     valign=MSO_ANCHOR.MIDDLE)

    # Speaker notes
    notes_text = (
        "Initial Assessment — Teacher Answers\n\n"
        "### Multiplication\n"
        "1. 43 × 6 = 258\n"
        "2. 68 × 4 = 272\n"
        "3. 476 × 82 = 39,032\n"
        "4. 6,038 × 591 = 3,568,458\n\n"
        "### Division\n"
        "1. 864 ÷ 4 = 216\n"
        "2. 132 ÷ 8 = 16 r 4 (or 16.5)\n"
        "3. 165 ÷ 6 = 27 r 3 (or 27.5)\n"
        "4. 615 ÷ 12 = 51.25 (or 51 r 3)\n\n"
        "### Application in Exam Questions\n"
        "1. 144 ÷ 6 = 24 cm\n"
        "2. 14 × 25 = 350 seats\n"
        "3. 145 ÷ 6 = 24 r 1 → 24 full boxes (1 egg left over)\n"
        "4. 130 ÷ 24 = 5 r 10 → 6 coaches needed; 6 × £75 = £450"
    )
    s.notes_slide.notes_text_frame.text = notes_text

    # Move newly created slide to position 0 (Slide 1)
    sldIdLst = prs.slides._sldIdLst
    new_slide_elm = sldIdLst[-1]
    sldIdLst.remove(new_slide_elm)
    sldIdLst.insert(0, new_slide_elm)

    return s


def main():
    prs = Presentation(PPTX_PATH)
    original_count = len(prs.slides)
    print(f'Original slide count: {original_count}')

    build_assessment_slide(prs)
    new_count = len(prs.slides)
    print(f'New slide count: {new_count}')
    assert new_count == original_count + 1

    prs.save(PPTX_PATH)
    print(f'Successfully updated {PPTX_PATH.name}')


if __name__ == '__main__':
    main()
