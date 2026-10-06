"""Purpose: Create new ratio versions with a first-slide assessment; preserve existing slides.

Main contents:
- Module-level setup and legacy command flow.

Used By: manual legacy command invocation.

Uses: scripts/build_ratio.py.

Libs: python-pptx (editable slides and deck inspection).

Legacy tooling: historical resource paths are retained; documentation changes do not authorize running it.
"""
raise SystemExit('Retired assessment migration: assessments are already in current decks; historical outputs are kept in Git.')
from pathlib import Path
from pptx import Presentation
import build_ratio as b
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'topics/ratio'
questions=[('Share £35 between A and B in the ratio 2:5. How much does A receive?','£10','7 parts = £35; 1 part = £5; A receives 2 × £5.'),('Red and blue counters are in the ratio 3:4. There are 12 red counters. How many blue counters are there?','16 blue counters','3 parts = 12; 1 part = 4; blue = 4 × 4.'),('Two ribbons have lengths in the ratio 2:7. The longer is 25 cm longer than the shorter. Find the shorter length.','10 cm','5 extra parts = 25 cm; 1 part = 5 cm; shorter = 2 × 5 cm.'),('Adults and children are in the ratio 5:8. There are 39 people altogether. How many more children than adults are there?','9 more children','13 parts = 39; 1 part = 3; difference = (8 − 5) × 3.')]
for answer,old,new in [(False,'Ratio_M04_questions_prev8.pptx','Ratio_M04_questions_prev9.pptx'),(True,'Ratio_M04_answers_prev5.pptx','Ratio_M04_answers_prev6.pptx')]:
 p=Presentation(OUT/old);s=b.base(p,'Ratio: initial assessment'+(' — answers' if answer else ''),'Match each answer to its question.' if answer else 'Answer all four questions. Show your working.','M04-IA'+('-A' if answer else ''))
 for i,(question,result,method) in enumerate(questions):
  x=.5+(i%2)*7.65;y=1.85+(i//2)*3.1;b.rect(s,x,y,7.35,2.85,'FFFFFF',b.LINE)
  b.text(s,x+.15,y+.15,7.05,1.4,f'{i+1}.  '+question,21,True)
  if answer:b.text(s,x+.15,y+1.6,7.05,1.12,'Answer: '+result+'\n'+method,19)
 p.slides._sldIdLst.insert(0,p.slides._sldIdLst[-1])
 temp=Path('/private/tmp')/new;p.save(temp);data=temp.read_bytes();(OUT/new).write_bytes(data);assert (OUT/new).read_bytes()==data
 print(new,len(p.slides),'slides')
