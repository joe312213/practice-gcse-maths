"""Purpose: Build M13 review decks from Markdown; keep practice answers separate.

Main contents:
- section
- table
- newdeck
- base
- line
- numberline
- question_grid
- assessment
- recap
- make

Used By: manual legacy command invocation.

Uses: scripts/build_ratio.py.

Libs: Pillow (font measurement/image rendering), python-pptx (editable slides and deck inspection).

Legacy tooling: historical resource paths are retained; documentation changes do not authorize running it.
"""
from pathlib import Path
import re,sys,hashlib
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from PIL import ImageFont
import build_ratio as b
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'content/M13_signed_addition_and_subtraction.md'
src=SRC.read_text();OUT=ROOT/'topics/signed_numbers';OUT.mkdir(exist_ok=True,parents=True)
def section(sn,answer=False):
 """Extract the requested numbered section from authored Markdown.

 Parameters: sn — one-based source slide number; answer — whether to render the answer variant.
 Used by: question_grid, make.
 """
 pat=rf'^## S{sn:02}'+('-A' if answer else '')+r' — [^\n]+\n(.*?)(?=^## |\Z)'
 return re.search(pat,src,re.M|re.S)[1]
def table(sec):
 """Read three-column practice rows from a Markdown section.

 Parameters: sec — Markdown section.
 Calls: b.splitcells.
 Used by: question_grid, make.
 """
 rows={}
 for line in sec.splitlines():
  if line.startswith('| '):
   cells=b.splitcells(line)
   if len(cells)==4 and cells[0] not in ('Step','Label','Q','---'):rows[cells[0]]=cells[1:]
 return rows
def newdeck():
 """Create a blank widescreen PowerPoint presentation.
 Used by: make.
 """
 p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9);return p
TITLE='Signed numbers: addition and subtraction'
def base(p,subtitle,ref):
 """Create a topic slide with its title, subtitle and stable reference.

 Parameters: p — presentation or drawing surface; subtitle — slide subtitle; ref — stable slide
 reference.
 Calls: b.base.
 Used by: question_grid, assessment, recap, make.
 """
 return b.base(p,TITLE,subtitle,ref)
def line(s,x,y,xx,yy,c=b.INK,arrow=False):
 """Draw a line between the supplied endpoints.

 Parameters: s — slide or source text as used by this helper; x — horizontal coordinate in inches; y
 — vertical coordinate in inches; xx — line end x coordinate in inches; yy — line end y coordinate
 in inches; c — coefficient or drawing colour as used here; arrow — whether the line needs an
 arrowhead.
 Calls: b.line.
 Used by: numberline.
 """
 b.line(s,x,y,xx,yy,c)
 sh=s.shapes[-1];sh.line.width=Pt(2 if arrow else 1)
 if arrow:sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(parse_xml('<a:tailEnd '+nsdecls('a')+' type="triangle"/>'))
def numberline(s,x,y,w,lo,hi,jumps,t):
 # Editable axis, ticks, coloured starts/endpoints and stepped jumps.
 """Draw labelled number positions and the specified signed jumps.

 Parameters: s — slide or source text as used by this helper; x — horizontal coordinate in inches; y
 — vertical coordinate in inches; w — width in inches; lo — lower number-line bound; hi — upper
 number-line bound; jumps — signed number-line jumps; t — challenge index or Bezier parameter as
 used here.
 Calls: line, xx, b.text.
 Used by: question_grid.

 Example in the caller's context: numberline(s, x, y, w, lo, hi, jumps, t)
 """
 def xx(n):
  """Map a number to its horizontal position on the number line.

  Parameters: n — number or numerator.
  Used by: numberline.
  """
  return x+.18+(n-lo)/(hi-lo)*(w-.36)
 line(s,xx(lo),y,xx(hi),y)
 for n in range(lo,hi+1):
  a=xx(n);line(s,a,y-.04,a,y+.04)
  label=b.text(s,a-.17,y+.07,.34,.22,str(n).replace('-','−'),10,n==0)
  label.text_frame.margin_left=label.text_frame.margin_right=0
  label.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
  if n<0:
   # Centre the digits on the tick; the minus sign extends to the left.
   font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',1000)
   label.left-=Pt(font.getlength('−')/100/2)
 for j,(start,end,label) in enumerate(jumps):
  yy=y-.34-j*.32
  line(s,xx(start),y-.06,xx(start),yy,b.COL[t]);line(s,xx(start),yy,xx(end),yy,b.COL[t],True);line(s,xx(end),yy,xx(end),y-.06,b.COL[t])
  b.text(s,x,y-.99+j*.29,w,.27,label,12,True,b.COL[t])
def question_grid(p,sn):
 """Build a worked or scaffolded signed-number question grid.

 Parameters: p — presentation or drawing surface; sn — one-based source slide number.
 Calls: table, section, base, b.text, b.headers, b.rect, numberline, b.progression.
 Used by: make.

 Example in the caller's context: question_grid(p, sn)
 """
 rows=table(section(sn));demo=sn==1
 s=base(p,'Worked demo • Follow each calculation.' if demo else f'Step-by-step practice {sn-1} • Show your working.',f'M13-S{sn:02}')
 if demo:
  b.text(s,.4,1.17,15.1,.6,'Adding a negative moves left: 4 + (−7) = 4 − 7 = −3.\nSubtracting a negative moves right: 6 − (−3) = 6 + 3 = 9.',16,True)
  left,cw=b.headers(s,True)
  # Shift headers below the method reminder.
  for sh in list(s.shapes)[-6:]:sh.top+=Inches(.65)
  y=2.36;heights=[.75,.85,.85,2.10,1.30]
 else:left,cw=b.headers(s,True);y=1.71;heights=[.9,1.38,1.38,1.38,1.36]
 for r,(label,vals) in enumerate(rows.items()):
  h=heights[r];b.rect(s,.35,y,left-.35,h,'E8EDF1',b.LINE);b.text(s,.41,y+.1,left-.47,h-.17,'The Full Question' if label=='Question' else label,17,True)
  for t,val in enumerate(vals):
   x=left+t*cw;b.rect(s,x,y,cw,h,b.TINT[t] if r==0 else 'FFFFFF',b.LINE)
   b.text(s,x+.08,y+.09,cw-.16,.60 if demo and label=='Calculate' else h-.16,val,24 if r==0 else 17,r==0)
   if demo and label=='Calculate':
    specs=[(-5,4,[(-4,0,'+4 then +3: total +7'),(0,3,'')]),(0,10,[(6,9,'subtract −3 = add 3')]),(-6,6,[(-5,3,'+8, then +2'),(3,5,'')])]
    lo,hi,jumps=specs[t];numberline(s,x+.12,y+1.66,cw-.24,lo,hi,jumps,t)
  y+=h
 if not demo:b.progression(s,True)
 s.notes_slide.notes_text_frame.text=section(sn)

def assessment(p,answers=False):
 """Append the topic initial-assessment slide.

 Parameters: p — presentation or drawing surface; answers — whether to show assessment answers.
 Calls: b.splitcells, base, b.rect, b.text.
 Used by: make.

 Example in the caller's context: assessment(p, answers)
 """
 sec=re.search(r'^## Initial assessment.*?\n(.*?)(?=^## |\Z)',src,re.S|re.M)[1]
 rows=[b.splitcells(l) for l in sec.splitlines() if re.match(r'^\| [1-4] \|',l)]
 s=base(p,'Initial assessment'+(' — answers' if answers else ' • Work out all four questions. Show your working.'),'M13-IA'+('-A' if answers else ''))
 for i,row in enumerate(rows):
  x=.5+(i%2)*7.65;y=1.85+(i//2)*3.05
  b.rect(s,x,y,7.35,2.75,'FFFFFF',b.LINE);b.text(s,x+.18,y+.2,6.9,.65,f'{row[0]}.  {row[2]}',30,True)
  if answers:b.text(s,x+.18,y+1.05,6.9,1.35,'Answer: '+row[3]+'\n'+row[4],23)
 s.notes_slide.notes_text_frame.text='Initial assessment; no hints or answers on question slide.'

def recap(p):
 """Append the topic rules/technique recap slide.

 Parameters: p — presentation or drawing surface.
 Calls: base, b.splitcells, b.rect, b.text.
 Used by: make.

 Example in the caller's context: recap(p)
 """
 s=base(p,'Rules recap • Start at the first number.','M13-RULES')
 block=re.search(r'^## Rules recap.*?\n(.*?)(?=^## |\Z)',src,re.M|re.S)[1]
 rules=[b.splitcells(l) for l in block.splitlines() if l.startswith('| ') and not l.startswith('| Rule') and not l.startswith('| ---')]
 for i,(rule,action,example) in enumerate(rules):
  x=.4+(i%2)*7.75;y=1.4+(i//2)*2.10
  b.rect(s,x,y,7.45,1.95,b.TINT[0] if i%2==0 else b.TINT[1],b.LINE)
  b.text(s,x+.15,y+.12,7.12,.43,rule,23,True,b.COL[0] if i%2==0 else b.COL[1])
  b.text(s,x+.15,y+.66,7.12,.55,action,20)
  b.text(s,x+.15,y+1.28,7.12,.48,example,25,True)
 b.text(s,.5,5.77,15,.70,'In 6 − (−3), the first − means subtract;\nthe − inside brackets belongs to the number −3.',22)
 b.rect(s,.4,6.65,15.2,.77,'E8EDF1',b.LINE)
 b.text(s,.55,6.78,14.9,.49,'To subtract a number, add its opposite. Keep the first number unchanged.',23,True)
 b.text(s,.5,7.68,15,.76,'Two negative numbers do not automatically give a positive answer:\n−4 + (−3) = −7.',22)
 s.notes_slide.notes_text_frame.text=block
 return s

def make():
 """Build the topic outputs from the authored inputs and save the legacy resources.
 Calls: newdeck, assessment, recap, question_grid, base, b.headers, table, section, b.rect, b.text,
 b.progression.

 Example in the caller's context: make()
 """
 q=newdeck();a=newdeck();assessment(q);recap(q);assessment(a,True)
 for sn in range(1,4):question_grid(q,sn)
 for sn in range(4,8):
  s=base(q,f'Independent practice {sn-3} • Choose a column. Show your working and check your answers.',f'M13-S{sn:02}');left,cw=b.headers(s)
  for k,vals in table(section(sn)).items():
   y=1.73+(int(k)-1)*1.10
   for t,v in enumerate(vals):
    x=left+t*cw;b.rect(s,x,y,cw-.025,1.1,b.TINT[t] if int(k)%2==0 else 'FFFFFF',b.LINE);b.text(s,x+.10,y+.26,cw-.23,.6,f'{k}.  {v}',26)
  b.progression(s)
 for sn in range(2,8):
  rows=table(section(sn,True));s=base(a,f'Answers • '+('Step-by-step' if sn<4 else 'Independent')+f' practice {sn-1 if sn<4 else sn-3}',f'M13-S{sn:02}-A');left,cw=b.headers(s)
  if sn<4:
   qs=table(section(sn))['Question']
   for t in range(3):
    x=left+t*cw;b.rect(s,x,1.8,cw-.03,6.5,b.TINT[t],b.LINE)
    b.text(s,x+.15,2,cw-.33,.6,qs[t],28,True)
    b.text(s,x+.15,3,cw-.33,.8,'Answer: '+rows['Answer'][t],28,True,b.COL[t])
    b.text(s,x+.15,4.1,cw-.33,3.6,rows['Method'][t],24)
  else:
   for k in range(1,7):
    for t,v in enumerate(rows[str(k)]):
     x=left+t*cw;y=1.76+(k-1)*1.03;b.rect(s,x,y,cw-.025,1.03,b.TINT[t] if k%2==0 else 'FFFFFF',b.LINE);b.text(s,x+.15,y+.24,cw-.32,.58,f'{k}.  {v}',28,True,b.COL[t])
  s=base(a,'Check your method • Correct the first wrong step.',f'M13-S{sn:02}-B');left,cw=b.headers(s)
  for t in range(3):
   for y,h,label in [(1.8,2.,'Method'),(3.93,2.47,'If you got…'),(6.53,1.85,'Check')]:
    x=left+t*cw;b.rect(s,x,y,cw-.035,h,b.TINT[t] if label!='If you got…' else 'FFFFFF',b.LINE);b.text(s,x+.1,y+.1,cw-.23,.4,label,20,True,b.COL[t]);b.text(s,x+.1,y+.58,cw-.23,h-.69,rows[label][t],19)
  s.notes_slide.notes_text_frame.text=section(sn,True)
 # Text-fit check for generated slides.
 for val,w,h,size,bold in b.checks:
  font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial'+(' Bold' if bold else '')+'.ttf',round(size*4));count=0
  for para in val.split('\n'):
   ln='';n=1
   for word in para.split():
    nxt=(ln+' '+word).strip()
    if font.getlength(nxt)>(w-.12)*72*4 and ln:n+=1;ln=word
    else:ln=nxt
   count+=n
  assert count*size*1.05<=(h-.05)*72+2,(val,w,h,size,count)
 for prs,name in [(q,'Signed_addition_subtraction_M13_questions_prev6.pptx'),(a,'Signed_addition_subtraction_M13_answers_prev1.pptx')]:
  temp=Path('/private/tmp')/name;prs.save(temp);payload=temp.read_bytes();(OUT/name).write_bytes(payload);assert (OUT/name).read_bytes()==payload;print(name,len(prs.slides),'slides')
if __name__=='__main__':
 raise SystemExit('Retired generator: use the current build commands in HANDOFF.md; historical outputs are kept in Git.')
if __name__=='__main__':make()
