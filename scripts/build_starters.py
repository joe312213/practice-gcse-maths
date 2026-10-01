"""Build the editable starter deck and its teacher answer key. Requires python-pptx."""
from pathlib import Path
from fractions import Fraction
import math
import hashlib
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from PIL import ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'GCSE_Maths_Revision_Starters_prev4.pptx'
KEY = ROOT / 'GCSE_Maths_Revision_Starters_answers_prev4.md'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
prs.core_properties.title = 'GCSE Maths Revision Starters'
prs.core_properties.subject = 'Lattice multiplication, bus stop division and application to exam problems'
prs.core_properties.author = 'Maths teaching resources'
BG, INK, MUTED, LINE = 'F7F8FA', '182B3A', '526472', 'D8E0E6'
COLORS = ['176B73', '3559A2', '754B87']
TINTS = ['EDF6F5', 'EFF3FA', 'F5F0F7']
LABELS = ['Start', 'Build', 'Confidence']
MARGIN, LABEL_W, FULL_W = .42, 2.30, 15.16
COL_W = (FULL_W - LABEL_W) / 3
GRID_X = MARGIN + LABEL_W
answer_sections = ['# Teacher answers — GCSE Maths Revision Starters v3',
                   'Answers also appear in the PowerPoint speaker notes. Worked diagrams are editable PowerPoint shapes.']
font_regular = '/System/Library/Fonts/Supplemental/Arial.ttf'
font_bold = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'
text_boxes = []


def box(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    effect=shape._element.find('.//' + qn('a:effectRef'))
    if effect is not None: effect.set('idx','0')
    shape.line.fill.background() if not line else None
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(.6)
    return shape


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.10)
    tf.margin_top = tf.margin_bottom = Inches(.04)
    tf.vertical_anchor = valign
    for i, line in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Arial'; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size*1.08)
    text_boxes.append((len(prs.slides), str(value), w, h, size, bold))
    return shape


def new_slide(module, title, subtitle, local):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    box(s, 0, 0, 16, .10, COLORS[module-1])
    text(s, MARGIN, .24, 13.6, .49, title, 29, bold=True)
    text(s, MARGIN, .78, 14.8, .30, subtitle, 14, MUTED)
    text(s, MARGIN, 8.62, 11, .23, 'GCSE Maths revision  •  Record your working in your booklet or webapp.', 11, MUTED)
    text(s, 12.5, 8.60, 3.05, .25, f'Module {module}  •  {local}/7  •  Slide {len(prs.slides)}', 11, MUTED)
    return s


def headers(s, grid=True, descriptors=None):
    start, width = (GRID_X, COL_W) if grid else (MARGIN, FULL_W/3)
    for c in range(3):
        x = start + c*width
        box(s, x, 1.23, width, .48, COLORS[c])
        text(s, x+.06, 1.26, width-.12, .37, LABELS[c], 20, 'FFFFFF', True)
        if descriptors:
            text(s, x+.04, 1.77, width-.08, .36, descriptors[c], 14, COLORS[c], True)


def grid(s, questions, row_labels, cells, heights, qheight=.65, sizes=None, diagrams=None):
    headers(s)
    y = 1.73
    box(s, MARGIN, y, LABEL_W, qheight, 'E8EDF1', LINE)
    text(s, MARGIN+.03, y+.04, LABEL_W-.06, qheight-.08, 'The Full Question', 16, bold=True, valign=MSO_ANCHOR.MIDDLE)
    for c,q in enumerate(questions):
        x = GRID_X + c*COL_W
        box(s, x, y, COL_W, qheight, TINTS[c], LINE)
        text(s, x+.06, y+.06, COL_W-.12, qheight-.12, q, 19 if qheight>1 else 26, bold=True, valign=MSO_ANCHOR.MIDDLE)
    y += qheight
    for r,(label,h) in enumerate(zip(row_labels,heights)):
        box(s, MARGIN, y, LABEL_W, h, 'E8EDF1', LINE)
        text(s, MARGIN+.07, y+.10, LABEL_W-.14, h-.20, f'{r+1}. {label}', 18, bold=True)
        for c in range(3):
            x = GRID_X + c*COL_W
            box(s,x,y,COL_W,h,'FFFFFF',LINE)
            diag = diagrams and r in diagrams
            diagram_h = diagrams[r][c][3] if diag else 0
            body_h = h-.18-diagram_h
            text(s,x+.07,y+.07,COL_W-.14,body_h,cells[r][c],(sizes or [19]*4)[r])
            if diag:
                kind,a,b,dh = diagrams[r][c]
                dy=y+h-dh
                if kind=='lattice': lattice(s,x+.12,dy,COL_W-.24,dh,a,b,c)
                elif kind=='bus': bus_stop(s,x+.10,dy,COL_W-.20,dh,a,b,c,c==2 and row_labels==DIVISION)
                elif kind=='both':
                    bus_stop(s,x+.08,dy,COL_W*.48,dh,145,9,c,False)
                    lattice(s,x+COL_W*.49,dy,COL_W*.49,dh,17,40,c)
        y += h
    assert y <= 8.48, y


def diagram_digit(group,x,y,w,h,value,size=18,color=INK,bold=False):
    shape=group.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=str(value); p.alignment=PP_ALIGN.CENTER
    p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold
    p.font.color.rgb=RGBColor.from_string(color)


def diagram_line(group,x1,y1,x2,y2,color=INK,width=1):
    line=group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    line.line.color.rgb=RGBColor.from_string(color); line.line.width=Pt(width)


def lattice(s,x,y,w,h,a,b,c):
    group=s.shapes.add_group_shape()
    group.name=f'Editable lattice {a} × {b}'
    aa,bb=str(a),str(b); n,m=len(aa),len(bb)
    cell=min(.62,(w-.74)/n,(h-.62)/m)
    gx=x+(w-n*cell)/2-.08; gy=y+.26
    size=18 if cell>=.49 else 15
    for j,ch in enumerate(aa): diagram_digit(group,gx+j*cell,gy-.28,cell,.25,ch,size,COLORS[c],True)
    for i,ch in enumerate(bb): diagram_digit(group,gx+n*cell+.03,gy+i*cell,.29,cell,ch,size,COLORS[c],True)
    sums=[0]*(n+m)
    for i,db in enumerate(bb):
        for j,da in enumerate(aa):
            xx,yy=gx+j*cell,gy+i*cell
            box(group,xx,yy,cell,cell,'FFFFFF',LINE)
            diagram_line(group,xx,yy+cell,xx+cell,yy,'8193A0',.8)
            tens,units=divmod(int(da)*int(db),10)
            diagram_digit(group,xx+.04*cell,yy+.02*cell,.48*cell,.47*cell,tens,size)
            diagram_digit(group,xx+.48*cell,yy+.49*cell,.48*cell,.47*cell,units,size)
            k=n-1-j+m-1-i; sums[k]+=units; sums[k+1]+=tens
    for j in range(n+1): diagram_line(group,gx+j*cell,gy,gx+j*cell,gy+m*cell,INK,.9)
    for i in range(m+1): diagram_line(group,gx,gy+i*cell,gx+n*cell,gy+i*cell,INK,.9)
    carry=0
    for k,v in enumerate(sums):
        total=v+carry; digit=total%10
        if k<n:
            xx,yy=gx+(n-1-k)*cell,gy+m*cell+.09
            diagram_digit(group,xx,yy,cell,.26,digit,size+1,COLORS[c],True)
            if carry: diagram_digit(group,xx+.02,gy+m*cell-.025,.16,.17,f'{carry}',10,'B14F00',True)
        else:
            row=m-1-(k-n); xx,yy=gx-.34,gy+row*cell
            leading=(k==len(sums)-1 and digit==0)
            diagram_digit(group,xx,yy,.29,cell,digit,size+1,'8193A0' if leading else COLORS[c],True)
            if carry: diagram_digit(group,gx-.18,yy+.01,.16,.17,f'{carry}',10,'B14F00',True)
        carry=total//10
    assert carry==0


def bus_stop(s,x,y,w,h,a,b,c,decimal=False):
    group=s.shapes.add_group_shape(); group.name=f'Editable bus stop {a} ÷ {b}'
    digits=list(str(a)); remainder=0; quot=[]; carries=[]
    for d in digits:
        carries.append(remainder); qv,remainder=divmod(remainder*10+int(d),b); quot.append(str(qv))
    if decimal and remainder:
        digits.append('.'); quot.append('.'); carries.append(0)
        while remainder:
            digits.append('0'); carries.append(remainder)
            qv,remainder=divmod(remainder*10,b); quot.append(str(qv))
    cell=min(.49,(w-.60)/len(digits)); width=cell*len(digits)
    gx=x+(w-width)/2+.12; gy=y+(h-.95)/2+.41
    size=24 if cell>.40 else 19
    diagram_line(group,gx-.08,gy,gx+width+.06,gy,INK,1.4)
    diagram_line(group,gx-.08,gy,gx-.08,gy+.42,INK,1.4)
    diagram_digit(group,gx-.52,gy+.02,.38,.38,b,size,COLORS[c],True)
    leading=True
    for i,d in enumerate(digits):
        xx=gx+i*cell
        diagram_digit(group,xx,gy+.02,cell,.38,d,size)
        qv=quot[i]
        if leading and qv=='0' and i<len(str(a))-1: qv=''
        else: leading=False
        diagram_digit(group,xx,gy-.41,cell,.35,qv,size,COLORS[c],True)
        if carries[i]:
            cw=.28 if carries[i]>=10 else .18
            diagram_digit(group,xx-cw*.55,gy-.01,cw,.18,carries[i],10 if carries[i]>=10 else 11,'B14F00',True)
    if remainder:
        diagram_digit(group,gx,gy+.40,width,.18,f'remainder {remainder}',12,COLORS[c])


def notes(s, heading, body):
    s.notes_slide.notes_text_frame.text = heading + '\n\n' + body
    answer_sections.append(f'\n## Slide {len(prs.slides)} — {heading}\n\n{body}')


def num(n):
    return f'{n:,}'


def division_answer(a,b,decimal=False):
    if decimal:
        f = Fraction(a,b)
        d = f.denominator
        for p in (2,5):
            while d%p == 0: d //= p
        assert d == 1, 'Choose a terminating decimal for this bank.'
        return f'{a/b:,.6f}'.rstrip('0').rstrip('.')
    q,r = divmod(a,b)
    return f'{q:,}' + (f' r {r}' if r else '')


LATTICE = ['Draw Grid & Arrange Digits','Find & Fill Cell Products',
           'Sum Diagonals & Track Carries','Write Final Answer']
DIVISION = ['Draw Frame & Position Digits','Divide Left-to-Right & Carry',
            'Track Remainders or Decimals','Write Final Quotient']
APPLICATION = ['Keywords & Calculation','Written Method','Ballpark Check & Math','Final Answer']

mul_sets = [
 [[(32,4),(61,8),(45,3),(72,6),(83,7),(95,9)],[(42,13),(26,34),(51,22),(63,41),(75,16),(82,57)],[(342,15),(612,24),(1234,125),(2105,304),(4316,1024),(7029,2316)]],
 [[(24,6),(57,4),(68,3),(39,7),(86,5),(74,8)],[(23,14),(46,25),(37,42),(58,36),(64,27),(79,53)],[(425,23),(708,46),(1326,243),(3048,512),(5607,1204),(8214,3067)]],
 [[(46,7),(29,8),(65,4),(78,6),(93,5),(87,9)],[(35,24),(62,17),(48,39),(76,28),(54,46),(89,65)],[(536,34),(809,27),(2047,368),(4158,427),(6309,2058),(9274,4163)]],
 [[(38,5),(56,9),(67,8),(49,6),(82,4),(97,7)],[(27,35),(43,26),(69,48),(85,37),(74,59),(96,68)],[(647,45),(902,38),(3156,489),(5027,596),(7418,3264),(8639,5076)]]
]
div_sets = [
 [[(48,4),(69,3),(84,2),(96,3),(55,5),(88,8)],[(72,3),(135,5),(172,4),(218,6),(341,7),(502,8)],[(514,4),(785,2),(1428,12),(3816,24),(4516,16),(8143,25)]],
 [[(46,2),(93,3),(66,6),(884,4),(555,5),(639,3)],[(84,3),(156,4),(235,6),(326,7),(458,8),(617,9)],[(627,4),(943,2),(2136,12),(4608,32),(5729,16),(9031,40)]],
 [[(82,2),(963,3),(848,4),(777,7),(888,8),(606,3)],[(95,4),(184,5),(267,8),(397,6),(523,7),(748,9)],[(738,4),(859,2),(3264,24),(5175,15),(6843,25),(7249,20)]],
 [[(62,2),(936,3),(448,4),(505,5),(666,6),(999,9)],[(108,4),(245,6),(314,7),(467,8),(586,9),(703,5)],[(849,4),(967,2),(4374,18),(6288,48),(7957,16),(8623,50)]]
]


def arithmetic_matrix(module, title, bank, index):
    desc = ['2 digits × 1 digit','2 digits × 2 digits','Up to 4 digits × 4 digits'] if module==1 else ['No carrying','Carrying; give remainders','Includes 2-digit divisors; give decimals']
    s = new_slide(module,title,f'Independent practice {index+1}  •  Choose a column. Show the full written method.',index+4)
    headers(s,False,desc)
    answers=[]
    for c, pairs in enumerate(bank):
        answers.append(f'### {LABELS[c]}')
        for r,(a,b) in enumerate(pairs):
            x,y,w = MARGIN+c*FULL_W/3, 2.21+r*1.015, FULL_W/3
            box(s,x,y,w,.995,'FFFFFF' if r%2==0 else TINTS[c],LINE)
            text(s,x+.12,y+.23,.50,.42,f'{r+1}.',19,COLORS[c],True)
            symbol = '×' if module==1 else '÷'
            q=f'{num(a)} {symbol} {num(b)}'
            text(s,x+.66,y+.20,w-.85,.55,q,30,bold=True)
            answer = num(a*b) if module==1 else division_answer(a,b,c==2)
            answers.append(f'{r+1}. {q} = **{answer}**')
    notes(s,f'{title} — independent practice {index+1}','\n\n'.join(answers))


def make_multiplication():
    s=new_slide(1,'Lattice Multiplication','Worked demo  •  Follow each step down your column. Small orange numbers show carries.',1)
    cells=[
      ['2 columns × 1 row.\nTop: 4, 3. Right: 6.', '2 columns × 2 rows.\nTop: 3, 4. Right: 1, 2.', '4 columns × 3 rows.\nTop: 4, 3, 1, 2. Right: 1, 5, 6.'],
      ['Cells: 24, 18.\nTens above; units below.', 'Top: 03, 04. Bottom: 06, 08.\nKeep the tens zeros.', 'Tens above; units below.'],
      ['Start at the bottom right:\n8 → write 8\n1 + 4 = 5 → write 5\n2 → write 2',
       'Start at the bottom right:\n8 → write 8\n4 + 0 + 6 = 10 → write 0, carry 1\n0 + 3 + 0 + 1 = 4 → write 4\n0 → leading zero',
       '2 → write 2\n1+0+6 = 7 → write 7\n1+2+0+5+8 = 16 → 6, carry 1\n0+0+1+1+5+4+1 = 12 → 2, carry 1\n0+1+3+2+0+1 = 7 → write 7\n0+2+4 = 6 → write 6\n0 → leading zero'],
      ['43 × 6 = 258','34 × 12 = 408','4,312 × 156 = 672,672']]
    # Seven diagonal sums in the last column; show each carry explicitly.
    grid(s,['43 × 6','34 × 12','4,312 × 156'],LATTICE,cells,[.80,2.65,1.85,.80],sizes=[17,17,15,21],
         diagrams={1:[('lattice',43,6,1.65),('lattice',34,12,1.65),('lattice',4312,156,2.15)]})
    notes(s,'Lattice Multiplication — worked demo',
          'Start: 43 × 6 = 258. Build: 34 × 12 = 408. Confidence: 4,312 × 156 = 672,672.\n\n'
          'The completed lattice diagrams are editable grouped PowerPoint shapes. Diagonals run from bottom left to top right. '
          'Place tens above each diagonal and units below it. Put answer digits at the ends of their diagonal tracks, '
          'down the left edge and along the bottom. Place carries at the bottom end of the NEXT diagonal in smaller text. '
          'Depth diagonal totals from right to left, including carries: 2, 7, 16, 12, 7, 6, 0. '
          'The last 1 in the fourth and fifth sums is a carry. Read the result in reverse order and omit the leading zero: 672672.')
    prompts=[
      'Draw the grid and diagonals. First number across the top; second number down the right edge.',
      'Multiply each row and column digit. Tens above; units below. Write a tens zero if the product is below 10.',
      'Start at the bottom right. Add each diagonal and any carry. Write its result digit at the bottom end. Put the next carry at the bottom end of the next diagonal in smaller text.',
      'Read down the left edge, then along the bottom. Omit leading zeros.']
    for i,qs in enumerate([[(53,7),(24,15),(1245,326)],[(84,5),(63,28),(3142,2045)]]):
        s=new_slide(1,'Lattice Multiplication',f'Step-by-step practice {i+1}  •  Complete each step in your booklet or webapp.',i+2)
        grid(s,[f'{num(a)} × {b}' for a,b in qs],LATTICE,[[p]*3 for p in prompts],[1.4,1.5,2.0,1.15],sizes=[19,20,19,20])
        notes(s,f'Lattice Multiplication — step-by-step practice {i+1}', '\n\n'.join(f'{LABELS[c]}: {num(a)} × {b} = **{num(a*b)}**' for c,(a,b) in enumerate(qs)))
    for i,bank in enumerate(mul_sets): arithmetic_matrix(1,'Lattice Multiplication',bank,i)


def make_division():
    s=new_slide(2,'Bus Stop Division','Worked demo  •  Start: exact answer. Build: remainder. Confidence: decimal answer.',1)
    cells=[
      ['Inside: 63.\nOutside on the left: 3.','Inside: 145.\nOutside on the left: 6.','Inside: 4,834.\nOutside on the left: 16.'],
      ['6 ÷ 3 = 2. Write 2 above 6.\n3 ÷ 3 = 1. Write 1 above 3.',
       '1 ÷ 6 = 0 r 1 → carry 1 to make 14.\n14 ÷ 6 = 2 r 2 → carry 2 to make 25.\n25 ÷ 6 = 4 r 1.',
       '4 ÷ 16 = 0 r 4 → make 48.\n48 ÷ 16 = 3; next digit is 3.\n3 ÷ 16 = 0 r 3 → make 34.\n34 ÷ 16 = 2 r 2.'],
      ['There is no final remainder.\nThe answer is a whole number.',
       'The final remainder is 1.\nWrite r 1 after the quotient.',
       'Add a decimal point and zeros.\n20 ÷ 16 = 1 r 4 → make 40.\n40 ÷ 16 = 2 r 8 → make 80.\n80 ÷ 16 = 5. No remainder.'],
      ['63 ÷ 3 = 21','145 ÷ 6 = 24 r 1','4,834 ÷ 16 = 302.125']]
    grid(s,['63 ÷ 3','145 ÷ 6','4,834 ÷ 16'],DIVISION,cells,[1.90,1.55,1.70,.85],sizes=[20,18,19,23],
         diagrams={0:[('bus',63,3,1.03),('bus',145,6,1.03),('bus',4834,16,1.03)]})
    notes(s,'Bus Stop Division — worked demo',
          'Start: 21. Build: 24 r 1. Confidence: 302.125.\n\n'
          'The editable bus stop diagrams show each quotient digit above its matching dividend digit. '
          'Show carried remainders as small prefixes. Preserve internal zeros, but omit an unnecessary leading zero in the final answer. '
          'For Depth, keep the zero in the tens position of 302.125. Align decimal points and show each carried remainder. Use multiples of 16 to check each quotient digit: 16, 32, 48, 64, 80. '
          'Checks: 21 × 3 = 63; 24 × 6 + 1 = 145; 302.125 × 16 = 4,834.')
    prompts=[
      'Draw the frame. Dividend inside; divisor outside on the left. Space the digits clearly.',
      'Work left to right. Write how many whole times the divisor fits above the matching digit. Carry the remainder as a small prefix to the next digit.',
      'Check the final remainder. Build: write any remainder as r. Confidence: add a decimal point and zeros. Continue until the remainder is 0.',
      'Write the complete answer. Keep place-value zeros. Align decimal points.']
    for i,qs in enumerate([[(88,4),(130,4),(5122,16)],[(66,3),(227,6),(3827,25)]]):
        s=new_slide(2,'Bus Stop Division',f'Step-by-step practice {i+1}  •  Build: give remainders. Confidence: give decimal answers.',i+2)
        grid(s,[f'{num(a)} ÷ {b}' for a,b in qs],DIVISION,[[p]*3 for p in prompts],[1.4,1.75,1.70,1.20],sizes=[19,20,19,20])
        notes(s,f'Bus Stop Division — step-by-step practice {i+1}', '\n\n'.join(f'{LABELS[c]}: {num(a)} ÷ {b} = **{division_answer(a,b,c==2)}**' for c,(a,b) in enumerate(qs)))
    for i,bank in enumerate(div_sets): arithmetic_matrix(2,'Bus Stop Division',bank,i)


# Question, independently calculable answer, teacher working. All contexts state the required constraint.
def q(text, answer, working): return (text, answer, working)

app_sets = [
 [
  [q('A library has 14 shelves with 32 books on each. How many books are there?',14*32,'14 × 32 = 448 books'),
   q('105 students form 7 equal groups. How many students are in each group?',105//7,'105 ÷ 7 = 15 students'),
   q('A shop has 26 crates with 41 bottles in each. How many bottles are there?',26*41,'26 × 41 = 1,066 bottles'),
   q('A 148 m rope is cut into 4 equal pieces. How long is each piece?',148//4,'148 ÷ 4 = 37 m'),
   q('A hall has 13 rows of 85 seats. How many seats are there?',13*85,'13 × 85 = 1,105 seats'),
   q('£175 is shared equally among 5 people. How much does each person receive?',175//5,'175 ÷ 5 = £35')],
  [q('130 eggs are packed in cartons of 6. How many full cartons can be packed?',130//6,'130 ÷ 6 = 21 r 4 → 21 full cartons'),
   q('32 passengers need cars with 5 passenger seats each. How many cars are needed?',math.ceil(32/5),'32 ÷ 5 = 6 r 2 → 7 cars'),
   q('250 chocolates are packed in boxes of 12. How many chocolates are left over?',250%12,'250 ÷ 12 = 20 r 10 → 10 chocolates'),
   q('A worker earns £17 per hour for 28 hours. What are their total wages?',17*28,'17 × 28 = £476'),
   q('142 rolls are sold in bags of 6. How many full bags can be made?',142//6,'142 ÷ 6 = 23 r 4 → 23 full bags'),
   q('A £132 bill is shared equally by 8 friends. How much does each pay?',132/8,'132 ÷ 8 = £16.50')],
  [q('45 tickets cost £16 each. A group gets £50 off the total. How much does it pay?',45*16-50,'45 × 16 = 720; 720 − 50 = £670'),
   q('115 passengers need vans with 8 passenger seats each. Hire costs £65 per van. Find the total cost.',math.ceil(115/8)*65,'115 ÷ 8 = 14 r 3 → 15 vans; 15 × 65 = £975'),
   q('A gym costs £22 per month plus a £30 joining fee. Find the cost for 14 months.',22*14+30,'22 × 14 = 308; 308 + 30 = £338'),
   q('430 cases fill boxes of 12. Each full box sells for £9. Find the income from full boxes.',(430//12)*9,'430 ÷ 12 = 35 r 10; 35 × 9 = £315'),
   q('A contractor charges £24 per hour for 15 hours, plus £45 for materials. Find the total bill.',24*15+45,'24 × 15 = 360; 360 + 45 = £405'),
   q('A shop sells 15 boxes of 12 headphones at £8 per headphone. Find the total sales income.',15*12*8,'15 × 12 = 180; 180 × 8 = £1,440')]
 ],
 [
  [q('18 trays each hold 24 plants. How many plants are there?',18*24,'18 × 24 = 432 plants'),
   q('156 pens are shared equally among 6 classes. How many pens does each class get?',156//6,'156 ÷ 6 = 26 pens'),
   q('A club orders 23 packs of 16 badges. How many badges does it order?',23*16,'23 × 16 = 368 badges'),
   q('£248 is shared equally among 8 people. How much does each person get?',248//8,'248 ÷ 8 = £31'),
   q('A runner covers 12 km each week for 27 weeks. How far do they run?',12*27,'12 × 27 = 324 km'),
   q('A 189 cm ribbon is cut into 7 equal pieces. Find the length of each piece.',189//7,'189 ÷ 7 = 27 cm')],
  [q('197 leaflets are bundled in groups of 8. How many complete bundles can be made?',197//8,'197 ÷ 8 = 24 r 5 → 24 bundles'),
   q('A camp has 83 people. Each tent sleeps 6. How many tents are needed?',math.ceil(83/6),'83 ÷ 6 = 13 r 5 → 14 tents'),
   q('A 275 cm strip is cut into 12 cm pieces. How many centimetres remain?',275%12,'275 ÷ 12 = 22 r 11 → 11 cm'),
   q('A printer produces 36 pages per minute. How many pages does it produce in 14 minutes?',36*14,'36 × 14 = 504 pages'),
   q('£154 is split equally among 4 workers. How much does each receive?',154/4,'154 ÷ 4 = £38.50'),
   q('253 photos go in albums holding 24 photos. How many albums are needed for all the photos?',math.ceil(253/24),'253 ÷ 24 = 10 r 13 → 11 albums')],
  [q('28 shirts cost £14 each. Delivery costs £18. Find the total cost.',28*14+18,'28 × 14 = 392; 392 + 18 = £410'),
   q('94 people need tables seating 6 each. Tables cost £12 each to hire. Find the total hire cost.',math.ceil(94/6)*12,'94 ÷ 6 = 15 r 4 → 16 tables; 16 × 12 = £192'),
   q('A club buys 24 packs of 15 drinks. It uses 87 drinks. How many remain?',24*15-87,'24 × 15 = 360; 360 − 87 = 273 drinks'),
   q('A baker has 286 biscuits. Full bags hold 8 and sell for £3 each. Find the sales income.',(286//8)*3,'286 ÷ 8 = 35 r 6; 35 × 3 = £105'),
   q('Four friends share the cost of 18 tickets at £12 each equally. How much does each friend pay?',18*12/4,'18 × 12 = 216; 216 ÷ 4 = £54'),
   q('A shop sells 17 cartons of 24 pens at £2 per pen. Find the total sales income.',17*24*2,'17 × 24 = 408; 408 × 2 = £816')]
 ],
 [
  [q('16 boxes each hold 35 cables. How many cables are there?',16*35,'16 × 35 = 560 cables'),
   q('252 stickers are shared equally among 7 students. How many does each student get?',252//7,'252 ÷ 7 = 36 stickers'),
   q('A cinema has 28 rows of 24 seats. Find the total number of seats.',28*24,'28 × 24 = 672 seats'),
   q('A 336 m cable is cut into 8 equal lengths. How long is each length?',336//8,'336 ÷ 8 = 42 m'),
   q('A charity buys 19 packs of 32 cards. How many cards does it buy?',19*32,'19 × 32 = 608 cards'),
   q('£315 is shared equally among 9 people. How much does each receive?',315//9,'315 ÷ 9 = £35')],
  [q('245 apples are packed in bags of 8. How many full bags can be made?',245//8,'245 ÷ 8 = 30 r 5 → 30 bags'),
   q('A hotel has 137 guests. Each shuttle carries 9 guests. How many shuttle trips are needed?',math.ceil(137/9),'137 ÷ 9 = 15 r 2 → 16 trips'),
   q('359 beads are used in bracelets of 12 beads. How many beads remain?',359%12,'359 ÷ 12 = 29 r 11 → 11 beads'),
   q('A technician earns £23 per hour for 16 hours. Find their total wages.',23*16,'23 × 16 = £368'),
   q('A £186 bill is shared equally by 8 people. How much does each person pay?',186/8,'186 ÷ 8 = £23.25'),
   q('Each floor tile covers 25 cm². What area do 48 tiles cover?',25*48,'25 × 48 = 1,200 cm²')],
  [q('A gardener charges £21 per hour for 18 hours, plus £64 for plants. Find the total bill.',21*18+64,'21 × 18 = 378; 378 + 64 = £442'),
   q('167 passengers need coaches with 24 passenger seats each. Hire costs £85 per coach. Find the total cost.',math.ceil(167/24)*85,'167 ÷ 24 = 6 r 23 → 7 coaches; 7 × 85 = £595'),
   q('A club buys 32 packs of 18 cards. It gives away 145 cards. How many remain?',32*18-145,'32 × 18 = 576; 576 − 145 = 431 cards'),
   q('397 oranges fill bags of 6. Each full bag sells for £4. Find the sales income.',(397//6)*4,'397 ÷ 6 = 66 r 1; 66 × 4 = £264'),
   q('Six people equally share the cost of 15 meals at £14 each. How much does each person pay?',15*14/6,'15 × 14 = 210; 210 ÷ 6 = £35'),
   q('A shop sells 23 boxes of 16 mugs at £5 per mug. Find the total sales income.',23*16*5,'23 × 16 = 368; 368 × 5 = £1,840')]
 ],
 [
  [q('27 packets each contain 18 labels. How many labels are there?',27*18,'27 × 18 = 486 labels'),
   q('384 books are shared equally among 8 shelves. How many books go on each shelf?',384//8,'384 ÷ 8 = 48 books'),
   q('A nursery has 34 trays of 16 seedlings. How many seedlings are there?',34*16,'34 × 16 = 544 seedlings'),
   q('£425 is shared equally among 5 people. How much does each receive?',425//5,'425 ÷ 5 = £85'),
   q('A hall has 26 rows of 35 seats. How many seats are there?',26*35,'26 × 35 = 910 seats'),
   q('A 294 cm strip is cut into 6 equal lengths. Find the length of each piece.',294//6,'294 ÷ 6 = 49 cm')],
  [q('367 bottles are packed in boxes of 12. How many complete boxes can be packed?',367//12,'367 ÷ 12 = 30 r 7 → 30 boxes'),
   q('A hostel has 158 guests. Each room holds 6 guests. How many rooms are needed?',math.ceil(158/6),'158 ÷ 6 = 26 r 2 → 27 rooms'),
   q('A 413 cm ribbon is cut into 15 cm lengths. How many centimetres remain?',413%15,'413 ÷ 15 = 27 r 8 → 8 cm'),
   q('A machine labels 28 bottles per minute. How many does it label in 19 minutes?',28*19,'28 × 19 = 532 bottles'),
   q('A £238 bill is shared equally by 8 friends. How much does each friend pay?',238/8,'238 ÷ 8 = £29.75'),
   q('319 files fit in folders holding 24 files. How many folders are needed for all the files?',math.ceil(319/24),'319 ÷ 24 = 13 r 7 → 14 folders')],
  [q('36 tickets cost £17 each. A group gets £75 off the total. How much does it pay?',36*17-75,'36 × 17 = 612; 612 − 75 = £537'),
   q('203 people need tables seating 8 each. Hire costs £15 per table. Find the total hire cost.',math.ceil(203/8)*15,'203 ÷ 8 = 25 r 3 → 26 tables; 26 × 15 = £390'),
   q('A shop buys 29 boxes of 24 notebooks. It sells 187 notebooks. How many remain?',29*24-187,'29 × 24 = 696; 696 − 187 = 509 notebooks'),
   q('A baker has 518 rolls. Full bags hold 12 and sell for £5 each. Find the sales income.',(518//12)*5,'518 ÷ 12 = 43 r 2; 43 × 5 = £215'),
   q('Seven people equally share the cost of 21 meals at £16 each. How much does each person pay?',21*16/7,'21 × 16 = 336; 336 ÷ 7 = £48'),
   q('A shop sells 28 boxes of 15 keyrings at £3 per keyring. Find the total sales income.',28*15*3,'28 × 15 = 420; 420 × 3 = £1,260')]
 ]
]


def make_application():
    s=new_slide(3,'Application to Exam Problems','Worked demo  •  Choose the calculation, use a written method, then check.',1)
    questions=['84 tins are shared equally among 4 boxes. How many tins go in each box?',
               '130 biscuits are packed in packs of 6. How many completely full packs can be made?',
               '145 students need minibuses with 9 passenger seats each. Hire costs £40 per minibus. Find the total cost.']
    cells=[
      ['“shared equally” among 4 boxes\n84 ÷ 4', '“packs of 6”; “completely full”\n130 ÷ 6', '“9 passenger seats”; “£40 per minibus”\n145 ÷ 9, then buses × 40'],
      ['Bus stop: 8 ÷ 4 = 2; 4 ÷ 4 = 1.\n84 ÷ 4 = 21.',
       'Bus stop: 13 ÷ 6 = 2 r 1;\n10 ÷ 6 = 1 r 4. So 130 ÷ 6 = 21 r 4.',
       'Bus stop: 14 ÷ 9 = 1 r 5;\n55 ÷ 9 = 6 r 1. So 16 r 1 → 17 buses.\nLattice: 17 × 40 = 680.'],
      ['80 ÷ 4 = 20.\n21 is close to 20.', '120 ÷ 6 = 20.\n21 full packs is reasonable.', '150 ÷ 10 ≈ 15; 15 × 40 = 600.\n£680 is reasonable: more buses are needed.'],
      ['21 tins per box','21 full packs; 4 biscuits left','£680 to hire 17 minibuses']]
    grid(s,questions,APPLICATION,cells,[1.05,2.45,1.05,.70],qheight=1.4,sizes=[17,17,17,20],
         diagrams={1:[('bus',84,4,1.15),('bus',130,6,1.15),('both',145,9,1.35)]})
    notes(s,'Application to Exam Problems — worked demo',
          'Start: 84 ÷ 4 = 21 tins per box. Build: 130 ÷ 6 = 21 r 4, so 21 full packs. '
          'Confidence: 145 ÷ 9 = 16 r 1, so 17 minibuses; 17 × 40 = £680.\n\n'
          'For Depth, 16 buses would seat only 144 students. The first answer must be interpreted before the next calculation. '
          'In the lattice for 17 × 40, top row cells are 04, 28; bottom row cells 00, 00. '
          'Diagonal totals from right: 0, 8, 6, 0 → 680. Both worked diagrams are provided as editable groups. '
          'Ballpark calculations check scale; they are not exact-answer proofs.')
    practices=[
      [q('42 packs each contain 15 pens. How many pens are there?',42*15,'42 × 15 = 630 pens. Estimate: 40 × 15 = 600.'),
       q('215 apples go into bags of 6. How many completely full bags can be made?',215//6,'215 ÷ 6 = 35 r 5 → 35 full bags. Estimate: 210 ÷ 6 = 35.'),
       q('A painter charges £18 per hour for 32 hours, plus £120 for paint. Find the total bill.',18*32+120,'18 × 32 = 576; 576 + 120 = £696. Estimate: 20 × 30 + 120 = 720.')],
      [q('168 students form 7 equal groups. How many students are in each group?',168//7,'168 ÷ 7 = 24 students. Estimate: 140 ÷ 7 = 20.'),
       q('A worker earns £16 per hour for 27 hours. What are their total wages?',16*27,'16 × 27 = £432. Estimate: 16 × 30 = 480.'),
       q('157 passengers need vans with 8 passenger seats each. Hire costs £45 per van. Find the total hire cost.',math.ceil(157/8)*45,'157 ÷ 8 = 19 r 5 → 20 vans; 20 × 45 = £900. Estimate: 160 ÷ 8 × 45 = 900.')]
    ]
    prompts=['Identify the keywords. Write down the calculation(s) you need to do.',
             'Set up and use the written method for your calculation(s).',
             'Check your answer using rough ballpark calculations. Does it make sense?',
             'Write your final answer to the question. Include the correct units.']
    for i,questions in enumerate(practices):
        s=new_slide(3,'Application to Exam Problems',f'Step-by-step practice {i+1}  •  Decide which calculation(s) each question needs.',i+2)
        grid(s,[v[0] for v in questions],APPLICATION,[[p]*3 for p in prompts],[1.3,1.3,1.3,1.3],qheight=1.4,sizes=[21]*4)
        notes(s,f'Application to Exam Problems — step-by-step practice {i+1}', '\n\n'.join(f'{LABELS[c]}: {v[2]}' for c,v in enumerate(questions)))
    for i,bank in enumerate(app_sets):
        s=new_slide(3,'Application to Exam Problems',f'Independent practice {i+1}  •  Choose a column. Show your working and check your answers.',i+4)
        headers(s,False)
        answers=[]
        for c,questions in enumerate(bank):
            answers.append(f'### {LABELS[c]}')
            for r,(question,answer,working) in enumerate(questions):
                x,y,w=MARGIN+c*FULL_W/3,1.76+r*1.095,FULL_W/3
                box(s,x,y,w,1.075,'FFFFFF' if r%2==0 else TINTS[c],LINE)
                text(s,x+.08,y+.12,.42,.38,f'{r+1}.',17,COLORS[c],True)
                text(s,x+.53,y+.09,w-.65,.90,question,18)
                answers.append(f'{r+1}. {working}')
        notes(s,f'Application to Exam Problems — independent practice {i+1}','\n\n'.join(answers))


def validate():
    assert len(prs.slides)==21
    for banks in (mul_sets,div_sets):
        pairs=[pair for bank in banks for col in bank for pair in col]
        assert len(pairs)==72 and len(set(pairs))==72
        assert all(len(col)==6 for bank in banks for col in bank)
    for bank in div_sets:
        for a,b in bank[0]:
            assert all(int(d)%b==0 for d in str(a)), (a,b)
        assert sum(1000 <= a <= 9999 and 10 <= b <= 99 for a,b in bank[2]) >= 4
    for bank in mul_sets:
        assert sum(1000 <= a <= 9999 and 100 <= b <= 999 for a,b in bank[2]) >= 2
        assert sum(1000 <= a <= 9999 and 1000 <= b <= 9999 for a,b in bank[2]) >= 2
    app_questions=[item[0] for bank in app_sets for col in bank for item in col]
    assert len(app_questions)==72 and len(set(app_questions))==72
    # Verify the manually described diagonal sums against the digit products.
    for a,b,expected in [(43,6,[8,5,2]),(34,12,[8,10,4,0]),(4312,156,[2,7,16,12,7,6,0])]:
        sums=[0]*(len(str(a))+len(str(b)))
        for i,da in enumerate(str(a)[::-1]):
            for j,db in enumerate(str(b)[::-1]):
                t,u=divmod(int(da)*int(db),10); sums[i+j]+=u; sums[i+j+1]+=t
        carry=0; totals=[]
        for v in sums:
            total=v+carry; totals.append(total); carry=total//10
        assert totals==expected,(a,b,totals)
    # Estimate line wrapping using actual Arial glyph widths before rendering.
    overflow=[]
    for slide,value,w,h,size,bold in text_boxes:
        font=ImageFont.truetype(font_bold if bold else font_regular,round(size*4))
        max_width=(w-.20)*72*4
        count=0
        for para in value.split('\n'):
            line=''; n=1
            for word in para.split():
                candidate=(line+' '+word).strip()
                if font.getlength(candidate)>max_width and line:
                    n+=1; line=word
                else: line=candidate
            count+=n
        needed=count*size*1.08
        available=(h-.08)*72
        if needed>available+2:
            overflow.append((slide,round(needed),round(available),value[:85]))
    if overflow:
        raise ValueError(f'Text needs more room: {overflow}')


if __name__=='__main__':
    make_multiplication(); make_division(); make_application(); validate()
    # Write to a local temporary file first, then verify the OneDrive copy.
    tmp=Path('/private/tmp/GCSE_Maths_Revision_Starters_prev4.pptx')
    prs.save(tmp)
    payload=tmp.read_bytes(); OUT.write_bytes(payload)
    assert hashlib.sha256(OUT.read_bytes()).digest()==hashlib.sha256(payload).digest()
    with zipfile.ZipFile(OUT) as archive: assert archive.testzip() is None
    KEY.write_text('\n\n'.join(answer_sections)+'\n')
    print(f'Created {OUT.name}: {len(prs.slides)} slides, {len(payload):,} bytes.')
    print(f'Created {KEY.name}. 216 independent questions; answers in all 21 slides’ notes.')
