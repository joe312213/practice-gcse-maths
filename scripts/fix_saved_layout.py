"""Apply approved progression arrows to plain grids and maintain the lattice templates."""
from copy import deepcopy
from pathlib import Path
import json
from pptx.util import Inches
import update_structure as u
ROOT=Path(__file__).resolve().parents[1]

def apply():
 registry=json.loads((ROOT/'content/topic_registry.json').read_text())
 ref=u.Presentation(ROOT/'topics/signed_numbers/Signed_addition_subtraction_M13_questions.pptx')
 reference=next(s for s in ref.slides if any('Step-by-step practice 2' in t for t in u.texts(s)))
 arrows=[deepcopy(sh._element) for sh in reference.shapes if sh.name.startswith('Progression curved arrow')]
 assert len(arrows)==2
 offsets=json.loads((ROOT/'content/progression_arrow.json').read_text())['arrows']
 count=0
 for entry in registry:
  path=ROOT/'topics'/entry['folder']/(entry['stem']+'_questions.pptx');p=u.Presentation(path);changed=False
  for s in p.slides:
   if not any('Independent practice' in t for t in u.texts(s)):continue
   for sh in list(s.shapes):
    if sh.name.startswith('Progression curved arrow'):s.shapes._spTree.remove(sh._element)
   headers=sorted([sh for sh in s.shapes if Inches(1.1)<sh.top<Inches(1.4) and Inches(.40)<sh.height<Inches(.65) and sh.width>Inches(3)],key=lambda sh:sh.left)
   assert len(headers)==3,(entry['id'],len(headers))
   cells=[sh for sh in s.shapes if sh.shape_type==1 and sh.top>Inches(1.6) and sh.height>Inches(.5) and sh.width>Inches(3) and sh.top+sh.height<Inches(8.6)]
   bottom=max((sh.top+sh.height for sh in cells),default=Inches(8.33))
   for i,xml in enumerate(arrows):
    el=deepcopy(xml);ident=el.find('.//'+u.qn('p:cNvPr'));ident.set('id',str(s.shapes._next_shape_id));ident.set('name',f'Progression curved arrow: Thread {i+1} to Thread {i+2}')
    off=el.find('.//'+u.qn('a:xfrm')).find(u.qn('a:off'));off.set('x',str(headers[i+1].left+Inches(offsets[i]['boundary_offset_inches'])));off.set('y',str(bottom+Inches(offsets[i]['bottom_offset_inches'])))
    s.shapes._spTree.insert_element_before(el,'p:extLst')
   count+=1;changed=True
  if entry['id']=='M01':
   for s in list(p.slides):
    if any('M01-TEMPLATES' in t for t in u.texts(s)):u.remove(p,s)
   target=next(s for s in p.slides if any('Step-by-step practice 1' in t for t in u.texts(s)))
   s=u.base(p,'Lattice multiplication — empty grids to copy','Size means columns × rows. Put the first number across the top and the second down the right.','M01-TEMPLATES')
   for n,m,x,y in [(2,2,1.25,1.60),(3,2,6.2,1.60),(3,3,11.3,1.60),(4,3,3.45,4.75),(4,4,9.6,4.75)]:
    u.b.text(s,x,y,3,.4,f'{n} × {m} grid',23,bold=True)
    gx=x+.25;gy=y+.6;c=.58
    group=s.shapes.add_group_shape();group.name=f'Empty {n} by {m} lattice template'
    for j in range(n+1):u.b.diagram_line(group,gx+j*c,gy,gx+j*c,gy+m*c)
    for i in range(m+1):u.b.diagram_line(group,gx,gy+i*c,gx+n*c,gy+i*c)
    for i in range(m):
     for j in range(n):u.b.diagram_line(group,gx+j*c,gy+(i+1)*c,gx+(j+1)*c,gy+i*c,'8193A0',.9)
   u.insert_before(p,s,target);changed=True
  if changed:u.save(p,path)
 print('Plain question grids with two approved arrows:',count)
if __name__=='__main__':
 apply()
 from compile_starters import compile_all
 compile_all()
