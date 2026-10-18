"""Purpose: Reconstruct M01/M02 historical import candidates from saved arithmetic sources and
decks. Candidate banks, catalogue changes and provenance are isolated for comparison; reviewed
live content remains authoritative.

Main contents:
- clean: Remove presentation markup from authored text; used by prepare_topic.
- operands: Parse integer ×/÷ text into left operand, operator and right operand; used by
  prepare_topic.
- answer_and_check: Return an answer and inverse check, using remainder notation at Build level.
- verify_answer: Check a saved correction against its operands; called by prepare_topic before
  accepting it.
- prepare_topic: Return a bank, catalogue pages, provenance and SVG text without writing files.
- main: Validate both reference topics, then write all candidates together into a fresh
  temporary directory.

Used By: optional manual reference reconstruction; never routine website builds.

Uses: historical content Markdown/JSON and saved decks/answer HTML; bank_candidates.py.

Libs: python-pptx for reading assessment examples; Python standard library for exact answers and
JSON.

Usage and output: From repository root: python3 scripts/prepare_web_arithmetic.py. Reads
historical Markdown/JSON/decks; write_candidates writes only a new temporary review directory.
"""
from pathlib import Path
from hashlib import sha256
from decimal import Decimal
import html
import json
import re
from pptx import Presentation
from bank_candidates import write_candidates

ROOT = Path(__file__).resolve().parents[1]
TOPICS = [
    ('M01', 'multiplication', 'Lattice_multiplication_M01', 'M01_lattice_multiplication.md', 'lattice'),
    ('M02', 'division', 'Bus_stop_division_M02', 'M02_bus_stop_division.md', 'bus'),
]


def clean(value):
    """Remove presentation markup from authored text; used by prepare_topic."""
    return html.unescape(re.sub('<[^>]+>', ' ', value)).strip().replace('**', '')


def operands(value):
    """Parse integer ×/÷ text into left operand, operator and right operand; used by prepare_topic."""
    left, operator, right = re.fullmatch(r'([\d,]+)\s*([×÷])\s*([\d,]+)', value).groups()
    return int(left.replace(',', '')), operator, int(right.replace(',', ''))


def answer_and_check(left, operator, right, level):
    """Return an answer and inverse check, using remainder notation at Build level.

    Called by prepare_topic.add with integer operands and zero-based level.
    Example: answer_and_check(13, '÷', 4, 1)  # ('3 r 1', '3 × 4 + 1 = 13')
    """
    if operator == '×':
        answer = str(left * right)
        return answer, f'{answer} ÷ {right} = {left}'
    quotient, remainder = divmod(left, right)
    if not remainder:
        answer = str(quotient)
    elif level == 1:
        answer = f'{quotient} r {remainder}'
    else:
        answer = str(Decimal(left) / Decimal(right)).rstrip('0').rstrip('.')
    check = f'{answer} × {right} = {left}' if '.' in answer else f'{quotient} × {right} + {remainder} = {left}'
    return answer, check


def verify_answer(item, expected):
    """Check a saved correction against its operands; called by prepare_topic before accepting it.

    Parameters: item contains a, b, method and id; expected is the saved answer string.
    Example: verify_answer({'a': 13, 'b': 4, 'method': 'bus', 'id': 'example'}, '3 r 1')
    """
    parts = expected.replace(',', '').split(' r ')
    if len(parts) == 2:
        valid = int(parts[0]) * item['b'] + int(parts[1]) == item['a']
    else:
        answer = Decimal(item['a']) * item['b'] if item['method'] == 'lattice' else Decimal(item['a']) / item['b']
        valid = Decimal(parts[0]) == answer
    assert valid, item['id']


def prepare_topic(topic, folder, stem, source, method, errors, support, assessments):
    """Return a bank, catalogue pages, provenance and SVG text without writing files.

    Parameters identify historical source files/method and loaded error/refinement/assessment data.
    Calls clean, operands, answer_and_check and verify_answer; used by main for TOPICS.
    Example: bank, pages, audit, diagrams = prepare_topic(*TOPICS[0], errors, support, assessments)
    """
    deck_path = ROOT / f'legacy/topics/{folder}/{stem}_questions.pptx'
    deck = Presentation(deck_path)
    slide_text = ['\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
    all_text = re.sub(r'\s|,', '', '\n'.join(slide_text))
    markdown = (ROOT / 'content' / source).read_text()
    saved_answers = (ROOT / f'legacy/topics/{folder}/{stem}_answers.html').read_text()
    questions, replacements = [], []
    diagrams = {}

    def add(question, identifier, kind, level, audit_source=True, **extra):
        """Append one historical candidate with its stable ID and any saved web refinements.

        Calls operands and answer_and_check; used by prepare_topic's section loops.
        audit_source checks historical provenance only; extra holds error/assessment fields.
        """
        if audit_source:
            assert re.sub(r'\s|,', '', question.split(' (')[0]) in all_text, (identifier, question)
        revised = support.get(f'maths:{identifier}', {}).get('q', question)
        if revised != question:
            replacements.append(dict(id=f'maths:{identifier}', source=question, website=revised))
        left, operator, right = operands(revised.split(' (')[0])
        answer, check = answer_and_check(left, operator, right, level)
        item = dict(id=f'maths:{identifier}', subject='maths', topic=topic, type=kind,
                    level=level, q=revised, a=left, b=right, answer=answer, check=check,
                    method=method, **extra)
        questions.append(item)
        return item

    # Saved examples are provenance inputs, not an authority over the live bank.
    assessment = re.findall(r'\d\.\s+([\d,]+\s*[×÷]\s*[\d,]+)', slide_text[0])
    assert len(assessment) == 4
    for index, question in enumerate(assessment):
        add(question, f'{topic}-IA-Q{index + 1}', 'assessment', [0, 1, 1, 2][index])
    sections = re.split(r'## Slide \d+ —[^\n]+', markdown)[1:]
    for number, section in enumerate(sections, 1):
        kind = 'demo' if number == 1 else 'scaffolded' if number < 4 else 'plain'
        rows = re.findall(r'^\| (?:The Full Question|[1-6]) \| (.+) \|$', section, re.M)
        for row, values in enumerate(rows, 1):
            for level, question in enumerate(values.split(' | ')):
                add(question, f'{topic}-S{number:02}-C{level + 1}-Q{row}', kind, level)

    # Retain original error diagrams and correction text inside the reference candidates.
    articles = re.findall(r'<article>(.*?)</article>', saved_answers.split(f'{topic}-SE')[1], re.S)[:9]
    assert len(articles) == 9
    corrections = [
        clean(re.search(r'<strong>Answer / Method:</strong>(.*?)</p>', article, re.S)[1])
        for article in articles
    ]
    for level, column in enumerate(errors[topic]):
        for index, error in enumerate(column):
            position = level * 3 + index
            question = error.get('question', f"{error.get('a', 0):,} × {error.get('b', 0):,}")
            identifier = f'{topic}-SE-C{level + 1}-Q{index + 1}'
            svg = re.search(r'<svg\b.*?</svg>', articles[position], re.S)[0]
            diagrams[f'website/static/data/working/{identifier}.svg'] = svg + '\n'
            options = [{'id': str(option), 'text': corrections[level * 3 + option]} for option in range(3)]
            item = add(question, identifier, 'errors', level,
                       errorImage=f'data/working/{identifier}.svg', errorOptions=options,
                       correctionId=str(index), correction=corrections[position],
                       wrong=error['incorrect_result'], error_note=corrections[position])
            verify_answer(item, error['correct_result'])
            item['answer'] = error['correct_result'].replace(',', '')
            item['check'] = error['check']
    for index, extra in enumerate(assessments[topic], 5):
        fields = {key: value for key, value in extra.items() if key not in ('q', 'level')}
        add(extra['q'], f'{topic}-IA-Q{index}', 'assessment', extra['level'], audit_source=False, **fields)
    for item in questions:
        item.update(support.get(item['id'], {}))

    method_rows = re.findall(r'^\| [1-4]\. (.*?) \| (.*?) \|', sections[0], re.M)
    prompts = [clean(text.replace('<br>', ' ')) for _, text in method_rows]
    steps = [
        clean(line.split(' | ')[1].replace('<br>', ' '))
        for line in sections[1].splitlines() if re.match(r'^\| [1-4]\.', line)
    ]
    aims = ['Use the lattice to keep each product in its place.'] * 3 if method == 'lattice' else [
        'Divide from left to right.',
        'Carry remainders; give any final remainder using r.',
        'Continue with decimal places until there is no remainder.',
    ]
    guidance = [{'aim': aim, 'steps': steps} for aim in aims]
    title = 'Lattice multiplication' if method == 'lattice' else 'Bus stop division'
    bank = dict(schema=1, subject='maths', topic=topic, title=title, method=method,
                rows=[label for label, _ in method_rows], prompts=prompts,
                questions=questions, teaching={'guidance': guidance})
    bank['revision'] = sha256(json.dumps(bank, sort_keys=True).encode()).hexdigest()[:16]
    pages = []
    for level in range(3):
        independent = [item for item in questions if item['type'] == 'plain' and item['level'] == level]
        slots = [[item['id'] for item in independent if f'-S{section:02}-' in item['id']] for section in range(4, 8)]
        pages.append(dict(type=0, level=level, slots=slots))
        pages.append(dict(type=1, level=level, slots=[
            [item['id'] for item in questions if item['type'] == 'errors' and item['level'] == level]
        ]))
    audit = dict(topic=topic, questions=len(questions), deck_sha256=sha256(deck_path.read_bytes()).hexdigest(),
                 answers_sha256=sha256(saved_answers.encode()).hexdigest(),
                 source_questions_found_in_saved_deck=True, web_replacements=replacements,
                 new_web_assessments=2, error_diagrams_preserved=9)
    return bank, pages, audit, diagrams


def main():
    """Validate both reference topics, then write all candidates together into a fresh temporary directory.

    Calls prepare_topic and write_candidates; no current bank, catalogue or audit is overwritten.
    """
    catalogue = json.loads((ROOT / 'website/src/lib/content/practice-pages.json').read_text())
    errors, support, assessments = [json.loads((ROOT / 'content' / name).read_text()) for name in (
        'spot_errors.json', 'M01_M02_web_support.json', 'web_assessments.json'
    )]
    files, audit = {}, []
    for definition in TOPICS:
        topic, folder, *_ = definition
        bank, pages, provenance, diagrams = prepare_topic(*definition, errors, support, assessments)
        files.update(diagrams)
        files[f'website/static/data/{folder}.json'] = json.dumps(bank, ensure_ascii=False, indent=2) + '\n'
        tags = next((item.get('tags', []) for item in catalogue['topics'] if item['bank'] == topic), [])
        catalogue['topics'] = [item for item in catalogue['topics'] if item['bank'] != topic]
        catalogue['topics'].append(dict(
            code=int(topic[1:]) - 1, bank=topic, title=bank['title'], pages=pages, tags=tags
        ))
        audit.append(provenance)
    catalogue['topics'].sort(key=lambda topic: topic['code'])
    files['website/src/lib/content/practice-pages.json'] = json.dumps(catalogue, ensure_ascii=False, indent=2) + '\n'
    files['docs/ARITHMETIC_IMPORT_AUDIT.json'] = json.dumps(audit, indent=2) + '\n'
    write_candidates(files)


if __name__ == '__main__':
    main()
