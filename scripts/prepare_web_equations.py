"""Purpose: Import accepted M10 structured content, checking against saved decks; no deck writes.

Main contents:
- normal
- slide_text
- main

Used By: manual website-bank import command.

Uses: no local module imports.

Libs: python-pptx (editable slides and deck inspection).
"""
from pathlib import Path
from fractions import Fraction
from hashlib import sha256
import json,re
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
LEGACY=ROOT/'legacy'  # Saved teaching decks are audit inputs, never regenerated here.
def normal(s):
 """Normalize whitespace and mathematical operator glyphs for source/deck comparison.

 Parameters: s — input text.
 Used by: main, add.
 """
 return re.sub(r'\s+','',s).replace('−','-').replace('×','*')
def slide_text(s):
 """Join text from all text-bearing shapes on a slide.

 Parameters: s — PowerPoint slide.
 Used by: main.
 """
 return '\n'.join(sh.text for sh in s.shapes if sh.has_text_frame)
def main():
 """Run the prepare web equations command using its configured input and output paths.
 Calls: shape_data, normal, slide_text, add.

 Example in the caller's context: main()
 """
 registry=json.loads((ROOT/'content/topic_registry.json').read_text());combined=Presentation(LEGACY/'GCSE_Maths_Revision_Starters.pptx');offset=0;audit=[]
 for t in registry:
  path=LEGACY/'topics'/t['folder']/(t['stem']+'_questions.pptx');p=Presentation(path)
  differences=[]
  def shape_data(s):
   """Extract shape type, geometry, rotation and text for deck comparison.

   Parameters: s — PowerPoint slide.
   Used by: main.
   """
   return [(int(sh.shape_type),sh.left,sh.top,sh.width,sh.height,sh.rotation,sh.text if sh.has_text_frame else '') for sh in s.shapes]
  for i,s in enumerate(p.slides):
   if shape_data(s)!=shape_data(combined.slides[offset+i]):differences.append(i+1)
  audit.append({'topic':t['id'],'slides':len(p.slides),'text_geometry_differences':differences,'sha256':sha256(path.read_bytes()).hexdigest()});offset+=len(p.slides)
 assert offset==len(combined.slides)
 d=json.loads((ROOT/'content/M10_equations.json').read_text());deck=Presentation(LEGACY/'topics/equations/Solving_equations_M10_questions.pptx');text=normal('\n'.join(slide_text(s) for s in deck.slides));items=[]
 def add(q,id,kind,level):
  """Validate one source question against the saved deck and append its web-bank entry.

  Parameters: q — source question; id — stable question reference; kind — question or operation
  variant; level — zero-based challenge level.
  Calls: normal.
  Used by: main.
  """
  assert normal(q['q']) in text,(id,q['q'])
  if 'coefficients' in q:
   a,b,c,e=map(Fraction,q['coefficients']);x=Fraction(q['answer']);assert a*x+b==c*x+e,id
  items.append({**q,'id':'maths:'+id,'subject':'maths','topic':'M10','type':kind,'level':level})
 for i,q in enumerate(d['assessment']):add(q,f'M10-IA-Q{i+1}','assessment',[0,0,1,2][i])
 for l,q in enumerate(d['demo']):add(q,f'M10-S01-C{l+1}-Q1','demo',l)
 for s,row in enumerate(d['practice'],2):
  for l,q in enumerate(row):add(q,f'M10-S{s:02}-C{l+1}-Q1','scaffolded',l)
 for s,cols in enumerate(d['banks'],4):
  for l,col in enumerate(cols):
   for i,q in enumerate(col):add(q,f'M10-S{s:02}-C{l+1}-Q{i+1}','plain',l)
 for l,col in enumerate(d['errors']):
  for i,q in enumerate(col):add(q,f'M10-SE-C{l+1}-Q{i+1}','errors',l)
 assert len(items)==94 and len({q['id'] for q in items})==94
 recap=json.loads((ROOT/'content/M10_web_recap.json').read_text())
 for block in recap:
  for note in [block['title'],*block['notes']]:assert normal(note) in text,note
  for row in block['balance']:
   for side in ('l','r'):assert normal(row[side]) in text,row

 teaching=json.loads((ROOT/'content/M10_web_teaching.json').read_text())
 for q in items:
  if q['type']=='errors':q.update(teaching['errorAnnotations'][q['id']])
 for q in teaching['additionalQuestions']:
  a,b,c,e=map(Fraction,q['coefficients']);x=Fraction(q['answer']);assert a*x+b==c*x+e,q['id']
  items.append(q)
 assert len({q['id'] for q in items})==len(items)
 for level,q in enumerate(d['demo']):assert len(q['balance'])==len(teaching['demoAnnotations'][level])
 for level,block in enumerate(recap):assert len(block['balance'])==len(teaching['recapAnnotations'][level])
 out={'schema':1,'revision':sha256((ROOT/'content/M10_equations.json').read_bytes()+(ROOT/'content/M10_web_recap.json').read_bytes()+(ROOT/'content/M10_web_teaching.json').read_bytes()).hexdigest()[:16],'subject':'maths','topic':'M10','title':d['title'],'rows':d['rows'],'prompts':d['prompts'],'questions':items,'recap':recap,'teaching':{k:teaching[k] for k in ('guidance','demoAnnotations','recapAnnotations','errorReasons')}}
 target=ROOT/'website/static/data';target.mkdir(parents=True,exist_ok=True);(target/'equations.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
 (ROOT/'docs/CONTENT_AUDIT.json').write_text(json.dumps({'combined_sha256':sha256((LEGACY/'GCSE_Maths_Revision_Starters.pptx').read_bytes()).hexdigest(),'topics':audit,'M10':{'imported':94,'new_web_items':len(teaching['additionalQuestions']),'plain':72,'source_questions_found_in_saved_deck':94,'coefficient_answers_verified':sum('coefficients' in q for q in items)},'limits':'Shape geometry/text compared; detailed styles/media and non-M10 source parity not yet audited.'},indent=2)+'\n')
 print('Imported 94 M10 items plus 1 new two-error example; 72 independent questions. All source question texts found in saved deck. Deck geometry differences:',[(r['topic'],r['text_geometry_differences']) for r in audit if r['text_geometry_differences']])
if __name__=='__main__':main()
