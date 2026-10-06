"""Purpose: Shared drawing primitives. Inches in; editable PPT shapes and/or standalone SVG out.

Main contents:
- Canvas

Used By: scripts/build_fraction_topic.py, scripts/rendering/bus_stop.py, scripts/rendering/equations.py, scripts/rendering/fractions.py, scripts/rendering/lattice.py, scripts/rendering/ratio.py, scripts/rendering/signed_numbers.py.

Uses: no local module imports.

Libs: Pillow (font measurement/image rendering), python-pptx (editable slides and deck inspection).

Legacy tooling: historical resource paths are retained; documentation changes do not authorize running it.
"""
import html,re
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR,MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from PIL import ImageFont
INK='182B3A'
class Canvas:
 def __init__(self,w,h,slide=None,x=0,y=0):
  """Initialize drawing geometry and the collected SVG output.

  Parameters: w — width in inches; h — height in inches; slide — editable PowerPoint slide; x —
  horizontal coordinate in inches; y — vertical coordinate in inches.
  """
  self.w=w;self.h=h;self.slide=slide;self.x=x;self.y=y;self.svg=[]
 def text(self,x,y,w,h,value,size=20,align='left',color=INK,bold=False):
  """Draw text with the supplied geometry and typography.

  Parameters: x — horizontal coordinate in inches; y — vertical coordinate in inches; w — width in
  inches; h — height in inches; value — text or numeric value to render; size — font size in points;
  align — text alignment; color — text/stroke colour hex string; bold — whether to use bold text.

  Example in the caller's context: canvas.text(x, y, w, h, value, size, align, color, bold)
  """
  value=str(value)
  if self.slide is not None:
   s=self.slide.shapes.add_textbox(Inches(x+self.x),Inches(y+self.y),Inches(w),Inches(h));tf=s.text_frame
   tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
   for i,line in enumerate(value.split('\n')):
    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align];p.space_after=Pt(0);p.line_spacing=Pt(size*1.12)
  for i,line in enumerate(value.split('\n')):
   tx=x+(w/2 if align=='center' else w if align=='right' else 0)
   self.svg.append(f'<text x="{tx*72}" y="{y*72+size*.91+i*size*1.12}" font-family="Arial,sans-serif" font-size="{size}" text-anchor="'+{'left':'start','center':'middle','right':'end'}[align]+f'" fill="#{color}" font-weight="{700 if bold else 400}">'+html.escape(line)+'</text>')
 def line(self,x,y,xx,yy,color=INK,width=1):
  """Draw a line between the supplied endpoints.

  Parameters: x — horizontal coordinate in inches; y — vertical coordinate in inches; xx — line end
  x coordinate in inches; yy — line end y coordinate in inches; color — text/stroke colour hex
  string; width — drawing width or stroke width as declared.
  """
  if self.slide is not None:
   sh=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(self.x+x),Inches(self.y+y),Inches(self.x+xx),Inches(self.y+yy));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width);sh._element.spPr.append(parse_xml('<a:effectLst '+nsdecls('a')+'/>'))
  self.svg.append(f'<line x1="{x*72}" y1="{y*72}" x2="{xx*72}" y2="{yy*72}" stroke="#{color}" stroke-width="{width}"/>')
 def rect(self,x,y,w,h,fill='FFFFFF',stroke=None):
  """Draw a rectangle with the requested fill and optional outline.

  Parameters: x — horizontal coordinate in inches; y — vertical coordinate in inches; w — width in
  inches; h — height in inches; fill — fill colour hex string; stroke — optional outline colour.
  """
  if self.slide is not None:
   sh=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(self.x+x),Inches(self.y+y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
   if stroke:sh.line.color.rgb=RGBColor.from_string(stroke)
   else:sh.line.fill.background()
   sh._element.spPr.append(parse_xml('<a:effectLst '+nsdecls('a')+'/>'))
  self.svg.append(f'<rect x="{x*72}" y="{y*72}" width="{w*72}" height="{h*72}" fill="#{fill}"'+(f' stroke="#{stroke}"' if stroke else '')+'/>')
 def math(self,x,y,value,size=24):
  """Draw text and stacked fraction tokens, returning their total horizontal width.

  Parameters: x — horizontal coordinate in inches; y — vertical coordinate in inches; value — text
  or numeric value to render; size — font size in points.
  Calls: self.text, self.line.

  Example in the caller's context: canvas.math(x, y, value, size)
  """
  font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',size*4);cursor=x
  for token in re.split(r'(\{[^{}]+/[^{}]+\})',value):
   if not token:continue
   if token.startswith('{'):
    num,den=token[1:-1].split('/');small=round(size*.78);sf=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',small*4)
    w=max(sf.getlength(num),sf.getlength(den))/(72*4)+.12
    self.text(cursor,y,w,.28,num,small,'center');self.line(cursor,y+.29,cursor+w,y+.29);self.text(cursor,y+.31,w,.28,den,small,'center');cursor+=w+.06
   else:
    w=font.getlength(token)/(72*4)+.03;self.text(cursor,y+.13,w,.38,token,size);cursor+=w
  return cursor-x
 def finish(self,label='Written method'):
  """Return the collected SVG with an accessible label.

  Parameters: label — accessible SVG label.
  """
  return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {self.w*72} {self.h*72}" role="img" aria-label="{html.escape(label,quote=True)}">'+''.join(self.svg)+'</svg>'
