"""Editable student-style working for SE cards; same diagrams embedded in HTML answers."""
import html
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import build_starters as b

class Paper:
 def __init__(self,slide,x,y):self.slide=slide;self.x=x;self.y=y;self.svg=[]
 def text(self,x,y,w,h,value,size=20,bold=False,align='left'):
  sh=self.slide.shapes.add_textbox(Inches(self.x+x),Inches(self.y+y),Inches(w),Inches(h))
  from PIL import ImageFont
  font=ImageFont.truetype(b.font_bold if bold else b.font_regular,round(size*4))
  for line in str(value).split('\n'):
   if font.getlength(line)>w*72*4+4:print('WORKING WIDTH:',value,w,size)
  tf=sh.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  for i,line in enumerate(str(value).split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
   p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(b.INK)
   p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.12)
   p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
   tx=(x+(w/2 if align=='center' else w if align=='right' else 0))*72
   self.svg.append(f'<text x="{tx:.2f}" y="{y*72+size*.91+i*size*1.12:.2f}" font-family="Arial,sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" text-anchor="'+{'left':'start','center':'middle','right':'end'}[align]+'">'+html.escape(line)+'</text>')
  return sh
 def line(self,x,y,xx,yy,width=1):
  b.diagram_line(self.slide,self.x+x,self.y+y,self.x+xx,self.y+yy,b.INK,width)
  self.svg.append(f'<line x1="{x*72}" y1="{y*72}" x2="{xx*72}" y2="{yy*72}" stroke="#182b3a" stroke-width="{width}"/>')
 def finish(self):return '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 360 152" role="img" aria-label="Student working with a deliberate mistake" style="width:100%;max-width:720px;background:white">'+''.join(self.svg)+'</svg>'

def bus(p,x,y,d,scale=1):
 """Digits occupy genuine place-value columns; remainders are small prefixes."""
 digits=d['digits'];q=d['quotient'];step=.46*scale;gx=x+.47*scale;gy=y+.36*scale
 p.line(gx-.06*scale,gy,gx+len(digits)*step+.03*scale,gy,1.3)
 p.line(gx-.06*scale,gy,gx-.06*scale,gy+.39*scale,1.3)
 p.text(x,gy+.025*scale,.36*scale,.35*scale,d['divisor'],22*scale,align='right')
 for i,ch in enumerate(digits):
  xx=gx+i*step
  p.text(xx,gy+.035*scale,step,.36*scale,ch,24*scale,align='center')
  if q[i]:
   qx=xx+d.get('decimal_shift',0)*step if q[i]=='.' else xx
   p.text(qx,gy-.36*scale,step,.34*scale,q[i],24*scale,align='center')
  rem=d.get('carries',{}).get(str(i))
  if rem is not None:p.text(xx-.06*scale,gy+.005*scale,.18*scale,.20*scale,rem,12*scale,align='center')
 if d.get('remainder') is not None:
  p.text(gx+len(digits)*step+.10*scale,gy-.33*scale,.8*scale,.30*scale,'r '+str(d['remainder']),18*scale)

def lattice(p,item):
 a,c=str(item['a']),str(item['b']);n,m=len(a),len(c);cell=.47;gx=.51;gy=.73
 for j,d in enumerate(a):p.text(gx+j*cell,gy-.29,cell,.29,d,20,align='center')
 for i,d in enumerate(c):p.text(gx+n*cell+.07,gy+i*cell+.12,.30,.29,d,20,align='center')
 for i,row in enumerate(item['_cells']):
  for j,v in enumerate(row):
   x=gx+j*cell;y=gy+i*cell
   p.line(x,y+cell,x+cell,y)
   p.text(x+.015,y+.015,.22,.22,v//10,16,align='center')
   p.text(x+.245,y+.245,.22,.22,v%10,16,align='center')
 for j in range(n+1):p.line(gx+j*cell,gy,gx+j*cell,gy+m*cell)
 for i in range(m+1):p.line(gx,gy+i*cell,gx+n*cell,gy+i*cell)
 for k,d in enumerate(item['_digits']):
  if k==n+m-1 and d=='0':continue
  if k<n:x=gx+(n-1-k)*cell;y=gy+m*cell+.09
  else:x=gx-.38;y=gy+(m-1-(k-n))*cell+.12
  p.text(x,y,.30,.30,d,21,align='center')
  carry=item['_incoming'][k]
  if carry:
   if k<n:cx=x-.06;cy=y-.13
   else:cx=x+.25;cy=y-.06
   big=item['_extra'] is not None and k==item['_extra'][0]+1
   if big:cy=y;cx=x+.23
   p.text(cx,cy,.19,.27,str(carry),21 if big else 11,align='center')
 p.text(2.5,.84,2.25,.4,'Answer',19)
 p.text(2.5,1.24,2.25,.42,f"{int(item['_result']):,}",25)

def equations(p,lines,x,y,width=4.6,size=21,step=.34):
 for i,line in enumerate(lines):p.text(x,y+i*step,width,.34,line,size)

def column(p,x,y,rows,width=.85,size=19):
 if rows[1].startswith('×'):
  carry=int(rows[0][-1])*int(rows[1][-1])//10
  if carry:p.text(x+width-.40,y-.13,.18,.19,carry,10,align='center')
 for i,row in enumerate(rows):
  p.text(x,y+i*.25,width,.28,row,size,align='right')
  if i in (1,len(rows)-2,len(rows)-1):p.line(x,y+(i+1)*.25-.005,x+width,y+(i+1)*.25-.005)

def draw(slide,x,y,mid,item):
 p=Paper(slide,x,y)
 if mid=='M01':lattice(p,item)
 elif mid=='M02':
  bus(p,.2,.69,item['bus'])
  p.text(.24,1.68,4.6,.35,item['final'],22)
 elif mid=='M03':
  scheme=item['layout']
  if scheme=='simple':equations(p,item['lines'],.25,.91,size=20,step=.33)
  elif scheme=='bus':
   bus(p,.2,.94,item['bus'],.9)
   p.text(.23,1.77,4.6,.30,item['final'],19)
  elif scheme=='bus_product':
   bus(p,.13,.95,item['bus'],.83)
   column(p,3.1,item.get('column_y',.82),item['column'],1.0,18)
   p.text(.23,1.78,2.7,.30,item['final'],18)
  elif scheme=='product_sum':
   column(p,.5,.82,item['column'],1.0,18)
   column(p,2.5,.91,item['sum'],1.0,18)
   p.text(2.2,1.78,2.7,.30,item['final'],18)
 elif mid=='M04':
  equations(p,item['left'],.23,.94,2.22,19,.31)
  equations(p,item['right'],2.57,.94,2.22,19,.31)
  p.text(.23,1.83,4.6,.28,item['final'],18)
 elif mid=='M13':
  equations(p,item['lines'],.35,.56,4.4,23,.38)
 item['_svg']=p.finish()
 return p
