"""Preserve saved teaching slides; refresh assessments/error activities and compile outputs.
Run from any directory: python3 practice/scripts/update_structure.py
Current topic files are required; restore missing accepted decks from Git.
"""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import json, re, html
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.opc.packuri import PackURI
import build_starters as b
import student_workings as student
import answer_layout
from rendering import estimates

ROOT=Path(__file__).resolve().parents[1]
DATA=json.loads((ROOT/'content/spot_errors.json').read_text())
TOPICS={
 'M01':('multiplication','Lattice_multiplication_M01','Lattice Multiplication'),
 'M02':('division','Bus_stop_division_M02','Bus Stop Division'),
 'M03':('problem_solving','Multiplication_division_problems_M03','Multiplication and Division Problems'),
 'M04':('ratio','Ratio_M04','Ratio: sharing and missing amounts'),
 'M13':('signed_numbers','Signed_addition_subtraction_M13','Signed numbers: addition and subtraction')}
ORDER=['M01','M02','M13','M03','M04']

def texts(s):return [sh.text for sh in s.shapes if sh.has_text_frame and sh.text]
def remove(p,s):
 for el in list(p.slides._sldIdLst):
  if p.part.related_part(el.rId)==s.part:
   p.part.drop_rel(el.rId);p.slides._sldIdLst.remove(el);return

def save(p,path):
 buf=BytesIO();p.save(buf);payload=buf.getvalue()
 path.write_bytes(payload)
 print(path.relative_to(ROOT),len(p.slides),'slides')

def save_text(path,value):
 if path.exists() and path.read_text()==value:return
 path.write_text(value)

def base(p,title,subtitle,ref):
 # Avoid python-pptx part-name collisions after deleting/reordering slides.
 for i,old in enumerate(p.slides):old.part._partname=PackURI(f"/ppt/slides/slide{1000+i}.xml")
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(b.BG)
 b.box(s,0,0,16,.1,b.COLORS[0]);b.text(s,.42,.24,15.1,.5,title,29,bold=True)
 b.text(s,.42,.81,15.1,.45,subtitle,17,b.MUTED)
 b.text(s,.42,8.60,15,.28,ref+'  •  Record your working in your booklet or webapp.',11,b.MUTED)
 return s

def assessment(p,mid,questions):
 s=base(p,TOPICS[mid][2]+' — initial assessment','Show your full written method. No calculators.',mid+'-IA')
 for i,q in enumerate(questions):
  x=.42+(i%2)*7.65;y=1.6+(i//2)*3.35
  b.box(s,x,y,7.5,3.12,'FFFFFF',b.LINE)
  b.text(s,x+.20,y+.22,7.1,2.7,f'{i+1}.  {q}',25 if mid=='M03' else 34,bold=True)
 if s.notes_slide.notes_text_frame is not None:s.notes_slide.notes_text_frame.text='Assessment answers are in the separate HTML answer file.'
 return s

def lattice_content(item):
 a,c=item['a'],item['b'];aa,cc=str(a),str(c);n,m=len(aa),len(cc)
 cells=[[int(x)*int(y) for x in aa] for y in cc]
 err=item['error'];explain=''
 if err=='cell':
  old=cells[0][0];cells[0][0]=old-10 if old>=10 else old+1
  explain=f'The top-left cell should be {int(aa[0])} × {int(cc[0])} = {old:02}, not {cells[0][0]:02}.'
 sums=[0]*(n+m)
 for i,row in enumerate(cells):
  for j,v in enumerate(row):
   k=n+m-2-i-j; t,u=divmod(v,10);sums[k]+=u;sums[k+1]+=t
 if err=='half':
  # Move the bottom-right tens into the units track.
  t=cells[-1][-1]//10;sums[0]+=t;sums[1]-=t
  explain=f'The {t} in the bottom-right cell belongs to the next diagonal (tens), not the units diagonal.'
 carry=0;digits=[];steps=[];dropped=False;extra=None;incoming_values=[]
 for k,v in enumerate(sums):
  incoming=carry
  if err=='drop' and carry and not dropped:
   incoming=0;dropped=True
  incoming_values.append(incoming)
  total=v+incoming;digit=total%10;carry=total//10
  if err=='carrydigit' and carry and extra is None:extra=(k,carry)
  steps.append(f'{v}'+(f'+{incoming}' if incoming else '')+f'={total}'+(f': write {digit}, carry {carry}' if carry else ''))
  digits.append(str(digit))
 result=''.join(reversed(digits)).lstrip('0') or '0'
 if err=='drop':explain='Include the carry from the previous diagonal in the next total.'
 if err=='carrydigit':
  k,v=extra
  # A large carry copied as an extra result digit at its diagonal boundary.
  result=result[:-k-1]+str(v)+result[-k-1:]
  explain='The carry is a small working digit, not an extra answer digit. Add it into the next diagonal once.'
 assert int(result) != a*c, (a,c,err,result)
 item.update(question=f'{a:,} × {c:,}',working='Cells: '+' / '.join(', '.join(f'{v:02}' for v in row) for row in cells)+'\nDiagonals from bottom right: '+'; '.join(steps)+'\nAnswer: '+result,correction=explain+f' Correct answer: {a*c:,}.',check=estimates.multiplication(a,c))
 item.update(_cells=cells,_digits=digits,_incoming=incoming_values,_extra=extra,_result=result)
 return cells,steps,result,extra

def errors(p,mid):
 s=base(p,TOPICS[mid][2]+' — spot the errors','Each student’s working contains a mistake. Find it, correct the working and give the right answer.',mid+'-SE')
 cw=15.16/3
 for t,col in enumerate(DATA[mid]):
  x=.42+t*cw;b.box(s,x,1.35,cw-.035,.46,b.COLORS[t]);b.text(s,x+.07,1.38,cw-.16,.38,b.LABELS[t],20,'FFFFFF',True)
  for i,item in enumerate(col):
   y=1.88+i*2.19;b.box(s,x,y,cw-.035,2.11,'FFFFFF',b.LINE)
   if mid=='M01':lattice_content(item)
   b.text(s,x+.10,y+.04,cw-.25,.92 if mid in ('M03','M04') else .44,f'{i+1}.  '+item['question'],18 if mid in ('M03','M04') else 21,bold=True)
   student.draw(s,x,y,mid,item)
 if s.notes_slide.notes_text_frame is not None:s.notes_slide.notes_text_frame.text='Deliberately incorrect student work. Corrections: see '+mid+'-SE in the HTML answers.'
 return s

def insert_before(p,s,target):
 el=p.slides._sldIdLst[-1];p.slides._sldIdLst.remove(el)
 for i,old in enumerate(p.slides):
  if old==target:p.slides._sldIdLst.insert(i,el);return
 raise ValueError('Target slide missing')

def clone(target,source):
 s=target.slides.add_slide(target.slide_layouts[6])
 for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
 for sh in source.shapes:s.shapes._spTree.insert_element_before(deepcopy(sh._element),'p:extLst')
 bg=source._element.cSld.find(qn('p:bg'))
 if bg is not None:s._element.cSld.insert(0,deepcopy(bg))
 # Saved decks use editable shapes only; refuse to silently lose media/charts.
 assert all(r.reltype.rsplit('/',1)[-1] in ('slideLayout','notesSlide') for r in source.part.rels.values())
 if source.has_notes_slide:
  ns=s.notes_slide._element
  ns.replace(ns.cSld,deepcopy(source.notes_slide._element.cSld))
 return s

# Small dependency-free Markdown renderer for the existing answer sources.
def inline(v):
 v=html.escape(v);v=v.replace('&lt;br&gt;','<br>')
 v=re.sub(r'\*\*(.*?)\*\*',r'<strong>\1</strong>',v)
 return re.sub(r'`(.*?)`',r'<code>\1</code>',v)
def md(value):
 out=[];table=False
 for line in value.splitlines():
  if line.startswith('|'):
   if re.match(r'^\|[\s:|\-]+$',line):continue
   if not table:out.append('<div class="table"><table>');table=True
   out.append('<tr>'+''.join('<td>'+inline(c.strip())+'</td>' for c in line.strip('|').split('|'))+'</tr>');continue
  if table:out.append('</table></div>');table=False
  if not line.strip():continue
  match=re.match(r'^(#{1,6}) (.*)',line)
  if match:
   level=min(5,len(match[1])+1);out.append(f'<h{level}>{inline(match[2])}</h{level}>')
  else:out.append('<p>'+inline(line)+'</p>')
 if table:out.append('</table></div>')
 return '\n'.join(out)

def error_answers(mid):
 out=[f'<h3>{mid}-SE — Spot the errors: corrections</h3><div class="three-columns">']
 for t,col in enumerate(DATA[mid]):
  out.append('<div><h4>'+b.LABELS[t]+'</h4>')
  for i,e in enumerate(col):
   out.append('<article><h4>'+str(i+1)+'. '+inline(e['question'])+'</h4><p><strong>If you got… (incorrect work)</strong><br>'+e.get('_svg',inline(e['working']).replace('\n','<br>'))+'</p><p><strong>Answer / Method:</strong> '+inline(e['correction'])+'</p>'+answer_layout.corrected(mid,e)+'<p><strong>Check:</strong> '+inline(e['check'])+'</p></article>')
  out.append('</div>')
 out.append('</div>')
 return '\n'.join(out)

IA={
'M01':['43 × 6 = 258','68 × 4 = 272','476 × 82 = 39,032','6,038 × 591 = 3,568,458'],
'M02':['864 ÷ 4 = 216','132 ÷ 8 = 16.5 or 16 r 4','165 ÷ 6 = 27.5 or 27 r 3','615 ÷ 12 = 51.25 or 51 r 3'],
'M03':['144 ÷ 6 = 24 cm','14 × 25 = 350 seats','145 ÷ 6 = 24 r 1: 24 full boxes, 1 egg left','130 ÷ 24 = 5 r 10: 6 coaches; 6 × £75 = £450']}

def answers(mid):
 out=[answer_layout.opening(mid,TOPICS[mid][2])+f'<h3>{mid}-IA — Initial assessment</h3>']
 if mid in IA:
  out.append(answer_layout.assessment_cell([f'<strong>Q{i+1}.</strong> '+inline(v) for i,v in enumerate(IA[mid])]))
 else:
  assessments=json.loads((ROOT/'content/assessment_answers.json').read_text())
  out.append(answer_layout.assessment_cell([f'<strong>Q{i+1}.</strong><p>'+inline(item['answer']).replace('\n','<br>')+'</p>' for i,item in enumerate(assessments[mid])]))
 if mid=='M13':
  source=(ROOT/'content/M13_signed_addition_and_subtraction.md').read_text()
  sections={int(n):body for n,body in re.findall(r'^## S(\d+)-A[^\n]*\n(.*?)(?=^## |\Z)',source,re.M|re.S)}
 else:
  source=(ROOT/f'content/answers/{mid}_answers.md').read_text()
  sections={int(n):body for n,body in re.findall(r'^## '+mid+r'-S(\d+)-A[^\n]*\n(.*?)(?=^## |\Z)',source,re.M|re.S)}
 for sn in range(2,8):
  if sn==4:out.append(error_answers(mid))
  out.append(f'<h3>{mid}-S{sn:02} — '+('Scaffolded practice '+str(sn-1) if sn<4 else 'Independent practice '+str(sn-3))+'</h3>'+answer_layout.worked(mid,sections[sn],md,sn))
 out.append('</div></section>');return '\n'.join(out)

def page(mids):
 nav=' · '.join(f'<a href="#{m}">{html.escape(TOPICS[m][2])}</a>' for m in mids)
 return '''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Maths starters — answers</title><style>'''+answer_layout.stylesheet()+'''</style><h1>Maths starters — answers</h1><p>Match the module, practice number, column and question. Correct the first wrong step, then redo the calculation. IA = initial assessment; SE = spot the errors. References stay fixed when slides move.</p><nav>'''+nav+'</nav>'+''.join(answers(m) for m in mids)+'</html>'

def main():
 from review_answer_patterns import review
 review(DATA)
 decks={}
 for mid,(folder,stem,title) in TOPICS.items():
  directory=ROOT/'topics'/folder;directory.mkdir(exist_ok=True)
  path=directory/(stem+'_questions.pptx')
  if not path.exists():
   raise FileNotFoundError(f'Restore the accepted topic deck from Git before building: {path}')
  p=Presentation(path)
  for s in list(p.slides):
   if any(mid+'-SE' in v for v in texts(s)):remove(p,s)
  target=next(s for s in p.slides if any('Independent practice 1' in v for v in texts(s)))
  insert_before(p,errors(p,mid),target)
  # Update only the obsolete physical-slide suffix in legacy footers.
  for idx,s in enumerate(p.slides,1):
   for sh in s.shapes:
    if sh.has_text_frame and re.match(r'Module [123]  •',sh.text):
     local=re.search(r'•  (\d)/7',sh.text)[1]
     sh.text_frame.paragraphs[0].runs[0].text=f'{mid}-S{int(local):02}  •  Topic slide {idx}'
  save(p,path);decks[mid]=p
  save_text(directory/(stem+'_answers.html'),page([mid]))
 from compile_starters import compile_all
 compile_all()
 # Human-readable exact content alongside the structured generator input.
 content=['# Spot the errors — exact slide content','Source: `spot_errors.json`; generated by `scripts/update_structure.py`. SE follows S03 and precedes S04; three questions per challenge column. All displayed workings below are deliberately incorrect. Corrections are shown only in HTML answers.']
 for mid in ORDER:
  content.append('\n## '+mid+' — '+TOPICS[mid][2])
  for t,col in enumerate(DATA[mid]):
   content.append('\n### '+b.LABELS[t])
   for i,e in enumerate(col):content.append('\n#### '+str(i+1)+'. '+e['question']+'\n\n'+e['working']+'\n\n**Written layout:** '+json.dumps({k:v for k,v in e.items() if k in ('bus','layout','column','sum','lines','left','right','final')},ensure_ascii=False)+'\n\n**Correction:** '+e['correction']+'\n\n**Check:** '+e['check'])
 (ROOT/'content/SPOT_ERRORS.md').write_text('\n'.join(content)+'\n')
 # A focused text-fit check for new text only; no broad existing-deck audit.
 from PIL import ImageFont
 warnings=[]
 for _,value,w,h,size,bold in b.text_boxes:
  font=ImageFont.truetype(b.font_bold if bold else b.font_regular,round(size*4));lines=0
  for para in value.split('\n'):
   current='';count=1
   for word in para.split():
    trial=(current+' '+word).strip()
    if font.getlength(trial)>(w-.20)*72*4 and current:count+=1;current=word
    else:current=trial
   lines+=count
  if lines*size*1.08>(h-.08)*72+2:warnings.append((value,lines,h,size))
 print('New text fit warnings:',warnings)
if __name__=='__main__':main()
