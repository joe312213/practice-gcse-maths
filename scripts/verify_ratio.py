"""Verify ratio v5 maths, variety, Markdown/slide agreement and diagnostic examples."""
from pathlib import Path
from fractions import Fraction
from collections import defaultdict, Counter
from math import gcd
import re
from pptx import Presentation
root=Path(__file__).resolve().parents[1]
src=(root/'content/M04_ratio.md').read_text();ans=(root/'content/answers/M04_answers.md').read_text()
def cells(line):return [x.strip() for x in line.strip('| ').split('|')]
qs={};aa={};teacher={}
for sn in range(4,8):
 section=src.split(f'### Slide {sn} —')[1].split(f'### Slide {sn+1} —')[0]
 qrows=[cells(x)[1:] for x in section.split('Teacher answers')[0].splitlines() if re.match(r'^\| [1-6] \|',x)]
 trows=[re.findall(r'\*\*(.*?)\*\*',x) for x in section.split('Teacher answers')[1].splitlines() if re.match(r'^\| [1-6] \|',x)]
 ac=re.search(r'^## M04-S'+f'{sn:02}'+r'-A[^\n]*\n(.*?)(?=^## M04-|\Z)',ans,re.M|re.S).group(1)
 arows=[cells(x)[1:] for x in ac.splitlines() if re.match(r'^\| [1-6] \|',x)]
 assert len(qrows)==len(arows)==len(trows)==6
 for k,(qr,ar,tr) in enumerate(zip(qrows,arows,trows),1):
  for t,(q,a,w) in enumerate(zip(qr,ar,tr),1):
   ref=f'M04-S{sn:02}-T{t}-Q{k}';qs[ref]=q;aa[ref]=a;teacher[ref]=w;assert a==w,(ref,a,w)
assert len(set(qs.values()))==72
variety=defaultdict(lambda:defaultdict(set));seen=set();ratio_counts=Counter();column_ratios=defaultdict(set);slide_ratios=defaultdict(Counter);difference_counts=Counter()
for line in (root/'content/M04_ratio_audit.md').read_text().splitlines():
 if not line.startswith('| M04-'):continue
 ref,ratio,k,kind,target,unit,context,names=cells(line);a,b=map(int,ratio.split(':'));k=Fraction(k);seen.add(ref)
 g=gcd(a,b);ratio_key=tuple(sorted((a//g,b//g)));ratio_counts[ratio_key]+=1
 column=ref.rsplit('-Q',1)[0]
 if kind=='difference':difference_counts[column]+=1
 assert ratio_key not in column_ratios[column],ref;column_ratios[column].add(ratio_key)
 slide_ratios[ref.split('-T')[0]][ratio_key]+=1
 q=qs[ref];assert re.search(r'ratio '+re.escape(ratio)+r'\b',q),ref
 # Audit values must be the actual numbers supplied in the question, not unrelated fixtures.
 q_without_ratio=re.sub(r'\b'+re.escape(ratio)+r'\b','',q)
 givens=re.findall(r'\d+(?:,\d{3})*(?:\.\d+)?',q_without_ratio)
 assert len(givens)==1 and Fraction(givens[0].replace(',',''))==k,(ref,q,givens,k)
 divisor={'total':a+b,'first':a,'second':b,'difference':abs(a-b)}[kind]
 p=k/divisor;x,y=p*a,p*b;assert x.denominator==y.denominator==1
 # Reconstruct both quantities and verify the stated relationship independently of the requested output.
 assert x*b==y*a
 assert {'total':x+y,'first':x,'second':y,'difference':abs(x-y)}[kind]==k
 expected={'a':[x],'b':[y],'total':[x+y],'difference':[abs(x-y)],'both':[x,y]}[target]
 actual=list(map(Fraction,re.findall(r'\d+(?:\.\d+)?',aa[ref])))
 assert actual==expected,(ref,actual,expected)
 assert unit in aa[ref],(ref,unit,aa[ref])
 if target=='both':assert all(n in aa[ref] for n in names.split(', ')),ref
 group=ref.split('-Q')[0];variety[group]['given'].add(kind);variety[group]['target'].add(target)
 category={'g':'mass','kg':'mass','£':'money','ml':'volume','L':'volume','m':'length','cm':'length','mm':'length','km':'length','min':'time'}.get(unit,'count')
 variety[group]['category'].add(category);variety[group]['context'].add(context)
assert seen==set(qs)
assert max(ratio_counts.values())<=3
assert all(max(v.values())<=2 for v in slide_ratios.values())
for ref,v in variety.items():
 assert len(v['category'])==6 and len(v['context'])==6,(ref,v)
 assert len(v['given'])>=3 and len(v['target'])>=2,(ref,v)
 if '-T1' in ref or '-T2' in ref:assert difference_counts[ref]==2,(ref,difference_counts[ref])
 if '-T3' in ref:assert {'total','difference'}<=v['given'] and 'difference' in v['target'],(ref,v)
qprs=Presentation(root/'topics/ratio/Ratio_M04_questions_prev8.pptx');aprs=Presentation(root/'topics/ratio/Ratio_M04_answers_prev5.pptx')
assert len(qprs.slides)==7 and len(aprs.slides)==12
for ref,q in qs.items():
 sn=int(ref.split('-')[1][1:]);text='\n'.join(sh.text for sh in qprs.slides[sn-1].shapes if sh.has_text_frame);assert q in text,(ref,q)
# Answer pages 5,7,9,11 carry the four independent answer banks; newlines separate paired answers.
for sn in range(4,8):
 slide=aprs.slides[4+2*(sn-4)];text=' '.join(sh.text for sh in slide.shapes if sh.has_text_frame)
 text=' '.join(text.split())
 for t in range(1,4):
  for k in range(1,7):
   val=aa[f'M04-S{sn:02}-T{t}-Q{k}'].replace('; ',' ');assert val in text,(sn,t,k,val)
# Explicitly check numerical consequences of the selected errors, including the new givens/objectives.
wrong_formulas={2:[12*4,35-(7-4),24*7],3:[8*5,18+(7-2),21*13],4:[12*5,15*4,10*7],5:[7*4,490/7,5*8],6:[70,30*4,12*9],7:[12*2,16*3,10*7]}
for sn,vals in wrong_formulas.items():
 sec=re.search(r'^## M04-S'+f'{sn:02}'+r'-A[^\n]*\n(.*?)(?=^## M04-|\Z)',ans,re.M|re.S).group(1)
 row=cells(re.search(r'^\| If you got… \|.*$',sec,re.M).group())[1:]
 for text,v in zip(row,vals):
  text=re.sub(r'^Q\d+ — ','',text)
  assert Fraction(re.search(r'\d[\d,]*(?:\.\d+)?',text).group().replace(',',''))==v,(sn,text,v)
print('Verified 72 distinct question texts, 72 answers and worked keys, 18 error examples, 7 question slides and 12 answer slides.')
print('All 12 independent columns mix six unit categories, at least three kinds of given information, and at least two objectives.')

print(f'{len(ratio_counts)} ratio relationships; none repeated in a column, at most twice per slide and three times overall, including reversals/equivalents.')

for sn,slide in enumerate(qprs.slides,1):
 assert not any(sh.has_text_frame and 'Try the next thread' in sh.text for sh in slide.shapes)
 arrows=[sh for sh in slide.shapes if sh.name.startswith('Progression curved arrow:')]
 assert len(arrows)==(0 if sn==1 else 2),(sn,len(arrows))
 for sh in arrows:
  assert sh.top>=int((8.03 if sn<4 else 8.25)*914400)-1 and sh.top+sh.height<int(8.72*914400)
for slide in aprs.slides:
 assert not any(sh.name.startswith('Progression curved arrow:') for sh in slide.shapes)
print('Two given-difference questions in each Guided/Core bank, plus two editable curved arrows on each of the six practice slides.')
